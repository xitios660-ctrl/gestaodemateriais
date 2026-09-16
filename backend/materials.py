"""Material catalogue and spreadsheet validation, shared by both database engines."""
import asyncio
import csv
import hashlib
import re
import unicodedata
import uuid
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from zipfile import ZipFile, BadZipFile

from bson import ObjectId
from fastapi import Depends, File, Form, HTTPException, UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field, model_validator
from starlette.responses import Response
from typing import Literal

MAX_ROWS = 5000
MAX_QUANTITY = 1_000_000_000
HEADERS = ["Categoria Pai", "Sub-item (Filho)", "Tipo de Medida (Unidade/Metro)", "Múltiplo"]
CATALOG_HEADERS = ["Sub-item", "Tipo de Medida (Unidade/Metro)", "Múltiplo"]


def name_key(value):
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).split()).casefold()


def plain(value):
    return "".join(c for c in unicodedata.normalize("NFKD", name_key(value)) if not unicodedata.combining(c))


class MaterialInput(BaseModel):
    name: str = Field(min_length=1, max_length=160)
    measure: Literal["unidade", "metro"] = "unidade"
    multiple: int = Field(default=1, strict=True, ge=1, le=MAX_QUANTITY)
    active: bool = True

    @model_validator(mode="after")
    def clean(self):
        self.name = " ".join(self.name.split())
        if not self.name:
            raise ValueError("Informe o nome do sub-item")
        if self.measure == "metro" and "multiple" not in self.model_fields_set:
            raise ValueError("Múltiplo é obrigatório para Metro")
        if self.measure == "unidade":
            self.multiple = 1
        return self


class MaterialBatchInput(BaseModel):
    items: list[MaterialInput] = Field(min_length=1, max_length=200)


def spreadsheet_rows(raw, filename):
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "csv":
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t|")
            rows = csv.reader(StringIO(content), dialect)
            result = []
            for row in rows:
                result.append(row)
                if len(result) > MAX_ROWS + 1:
                    raise ValueError(f"Use até {MAX_ROWS} linhas por planilha")
            return result
        except csv.Error:
            raise ValueError("CSV inválido. Separe as quatro colunas por vírgula ou ponto e vírgula")
    if ext != "xlsx":
        raise ValueError("Envie um arquivo .xlsx ou .csv")
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError("A planilha descompactada excede o limite de 40 MB")
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.worksheets[0]
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > 50:
                raise ValueError(f"Use até {MAX_ROWS} linhas e somente as quatro colunas do modelo")
            rows = []
            for cells in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"Linha {cells[0].row}: substitua fórmulas pelos valores")
                rows.append([cell.value for cell in cells])
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError(f"Use até {MAX_ROWS} linhas por planilha")
            return rows
        finally:
            workbook.close()
    except (BadZipFile, KeyError, OSError):
        raise ValueError("Não foi possível ler o Excel. Salve novamente no formato .xlsx")


def parse_spreadsheet(raw, filename):
    rows = spreadsheet_rows(raw, filename)
    if not rows:
        raise ValueError("A planilha está vazia")
    aliases = [
        {"categoria pai", "categoria"},
        {"sub-item (filho)", "sub-item", "subitem", "subitem (filho)", "filho"},
        {"tipo de medida (unidade/metro)", "tipo de medida", "medida"},
        {"multiplo", "multiplo (obrigatorio se for metro)"},
    ]
    headers = [plain(value) for value in rows[0]]
    indexes = [next((i for i, header in enumerate(headers) if header in choices), None) for choices in aliases]
    if any(index is None for index in indexes) or any(h and not any(h in a for a in aliases) for h in headers) or len([h for h in headers if h]) != 4:
        raise ValueError("Colunas esperadas: " + " | ".join(HEADERS))
    parsed, errors, seen = [], [], {}
    for number, row in enumerate(rows[1:], 2):
        if not any(str(value or "").strip() for value in row):
            continue
        values = [row[i] if i < len(row) else None for i in indexes]
        parent, child, measure, multiple = values
        try:
            if any(str(v or "").lstrip().startswith("=") for v in values):
                raise ValueError("Substitua fórmulas pelos valores")
            parent = " ".join(str(parent or "").split())
            if not parent or len(parent) > 120:
                raise ValueError("Informe uma categoria pai com até 120 caracteres")
            unit = plain(measure)
            if unit not in {"unidade", "unidades", "un", "metro", "metros", "m"}:
                raise ValueError("Tipo de medida deve ser Unidade ou Metro")
            unit = "metro" if unit in {"metro", "metros", "m"} else "unidade"
            if unit == "metro" and (multiple is None or str(multiple).strip() == ""):
                raise ValueError("Múltiplo é obrigatório para Metro")
            factor = Decimal(str(multiple).replace(",", ".")) if multiple not in (None, "") else Decimal(1)
            if not factor.is_finite() or factor != factor.to_integral_value() or not 1 <= factor <= MAX_QUANTITY:
                raise ValueError("Múltiplo deve ser um número inteiro positivo")
            item = MaterialInput(name=str(child or ""), measure=unit, multiple=int(factor)).model_dump()
            key = (name_key(parent), name_key(item["name"]))
            if key in seen and (seen[key]["measure"], seen[key]["multiple"]) != (item["measure"], item["multiple"]):
                raise ValueError("O mesmo sub-item aparece com regras diferentes na planilha")
            parsed.append({"row": number, "category": parent, **item})
            seen[key] = item
        except (ValueError, InvalidOperation) as exc:
            message = str(exc)
            if hasattr(exc, "errors"):
                message = exc.errors()[0]["msg"]
            errors.append({"row": number, "message": message})
    if not parsed and not errors:
        raise ValueError("Inclua pelo menos um sub-item abaixo do cabeçalho")
    return parsed, errors



def parse_catalog_spreadsheet(raw, filename):
    """Parse a spreadsheet for one already-created catalog. It never creates categories."""
    rows = spreadsheet_rows(raw, filename)
    if not rows:
        raise ValueError("A planilha está vazia")

    aliases = [
        {"sub-item", "sub-item (filho)", "subitem", "filho", "item"},
        {"tipo de medida (unidade/metro)", "tipo de medida", "medida"},
        {"multiplo", "multiplo (obrigatorio se for metro)"},
    ]
    headers = [plain(value) for value in rows[0]]
    indexes = [next((i for i, header in enumerate(headers) if header in choices), None) for choices in aliases]
    if any(index is None for index in indexes) or len([h for h in headers if h]) != 3:
        raise ValueError("Colunas esperadas: " + " | ".join(CATALOG_HEADERS))

    parsed, errors, seen = [], [], {}
    for number, row in enumerate(rows[1:], 2):
        if not any(str(value or "").strip() for value in row):
            continue
        values = [row[i] if i < len(row) else None for i in indexes]
        child, measure, multiple = values
        try:
            if any(str(v or "").lstrip().startswith("=") for v in values):
                raise ValueError("Substitua fórmulas pelos valores")
            unit = plain(measure)
            if unit not in {"unidade", "unidades", "un", "metro", "metros", "m"}:
                raise ValueError("Tipo de medida deve ser Unidade ou Metro")
            unit = "metro" if unit in {"metro", "metros", "m"} else "unidade"
            if unit == "metro" and (multiple is None or str(multiple).strip() == ""):
                raise ValueError("Múltiplo é obrigatório para Metro")
            factor = Decimal(str(multiple).replace(",", ".")) if multiple not in (None, "") else Decimal(1)
            if not factor.is_finite() or factor != factor.to_integral_value() or not 1 <= factor <= MAX_QUANTITY:
                raise ValueError("Múltiplo deve ser um número inteiro positivo")
            item = MaterialInput(name=str(child or ""), measure=unit, multiple=int(factor)).model_dump()
            key = name_key(item["name"])
            if key in seen and (seen[key]["measure"], seen[key]["multiple"]) != (item["measure"], item["multiple"]):
                raise ValueError("O mesmo sub-item aparece com regras diferentes na planilha")
            parsed.append({"row": number, **item})
            seen[key] = item
        except (ValueError, InvalidOperation) as exc:
            message = str(exc)
            if hasattr(exc, "errors"):
                message = exc.errors()[0]["msg"]
            errors.append({"row": number, "message": message})
    if not parsed and not errors:
        raise ValueError("Inclua pelo menos um sub-item abaixo do cabeçalho")
    return parsed, errors


def serialize_catalog(doc):
    data = dict(doc)
    data["id"] = str(data.pop("_id"))
    data["materials"] = list(data.get("materials", []))
    return data


def public_catalog(doc):
    data = serialize_catalog(doc)
    data["materials"] = [item for item in data.get("materials", []) if item.get("active", True)]
    return data


async def save_catalog_materials(database, catalog, items):
    query = {
        "_id": catalog["_id"],
        "materials": catalog["materials"] if "materials" in catalog else {"$exists": False},
    }
    result = await database.material_catalogs.update_one(
        query,
        {"$set": {"materials": items, "updated_at": datetime.now(timezone.utc).isoformat()}},
    )
    if not result.matched_count:
        raise HTTPException(
            409,
            "O catálogo mudou durante a operação. Atualize e tente novamente",
        )


def plan_catalog_items(rows, catalog):
    items = list(catalog.get("materials", []))
    by_name = {name_key(item["name"]): item for item in items}
    preview, errors = [], []
    counts = {"items_created": 0, "skipped": 0}

    for row in rows:
        existing = by_name.get(name_key(row["name"]))
        action = "criar"
        if existing:
            if existing["measure"] != row["measure"] or existing["multiple"] != row["multiple"]:
                errors.append({
                    "row": row["row"],
                    "message": "Sub-item já cadastrado com outra medida ou múltiplo. Edite-o manualmente",
                })
                continue
            counts["skipped"] += 1
            action = "manter"
        else:
            if len(items) >= MAX_ROWS:
                errors.append({"row": row["row"], "message": f"Limite de {MAX_ROWS} sub-itens por catálogo"})
                continue
            item = {k: row[k] for k in ("name", "measure", "multiple", "active")}
            item["id"] = str(uuid.uuid4())
            items.append(item)
            by_name[name_key(item["name"])] = item
            counts["items_created"] += 1
        preview.append({**row, "action": action})

    return items, preview, counts, errors

def catalog_plan(rows, categories):
    by_name, errors = {}, []
    for category in categories:
        key = name_key(category["name"])
        if key in by_name:
            by_name[key] = None
        else:
            by_name[key] = category
    planned, preview, counts = {}, [], {"categories_created": 0, "items_created": 0, "skipped": 0}
    for row in rows:
        key = name_key(row["category"])
        if key in by_name and by_name[key] is None:
            errors.append({"row": row["row"], "message": "Há categorias com este mesmo nome. Renomeie-as antes de importar"})
            continue
        if key not in planned:
            current = by_name.get(key)
            if current is None:
                current = {"_id": ObjectId(hashlib.sha256(key.encode()).hexdigest()[:24]), "name": row["category"], "icon": "PackagePlus", "description": "", "lead_time_hours": 24, "owners": [], "fields": [], "template_columns": [], "active": True, "created_at": datetime.now(timezone.utc).isoformat()}
                counts["categories_created"] += 1
            planned[key] = {"category": current, "new": key not in by_name, "materials": list(current.get("materials", [])), "by_name": {name_key(item["name"]): item for item in current.get("materials", [])}}
        plan = planned[key]
        existing = plan["by_name"].get(name_key(row["name"]))
        action = "criar"
        if existing:
            if existing["measure"] != row["measure"] or existing["multiple"] != row["multiple"]:
                errors.append({"row": row["row"], "message": "Sub-item já cadastrado com outra medida ou múltiplo. Edite-o manualmente"})
                continue
            counts["skipped"] += 1
            action = "manter"
        else:
            if len(plan["materials"]) >= MAX_ROWS:
                errors.append({"row": row["row"], "message": f"Limite de {MAX_ROWS} sub-itens por categoria"})
                continue
            item = {k: row[k] for k in ("name", "measure", "multiple", "active")}
            item["id"] = str(uuid.uuid5(uuid.NAMESPACE_URL, f"{plan['category']['_id']}/{name_key(row['name'])}"))
            plan["materials"].append(item)
            plan["by_name"][name_key(item["name"])] = item
            counts["items_created"] += 1
        preview.append({**row, "action": action, "category_active": plan["category"].get("active", True)})
    return list(planned.values()), preview, counts, errors


async def save_materials(database, category, items):
    query = {"_id": category["_id"], "materials": category["materials"] if "materials" in category else {"$exists": False}}
    result = await database.categories.update_one(query, {"$set": {"materials": items}})
    if not result.matched_count:
        raise HTTPException(409, "O catálogo mudou durante a operação. Atualize e tente novamente; itens já importados serão reconhecidos")


def register_material_routes(router, get_db, require_admin, audit):
    @router.get("/material-catalogs")
    async def list_material_catalogs():
        database = get_db()
        docs = await database.material_catalogs.find({}).sort("created_at", 1).to_list(500)
        return [public_catalog(doc) for doc in docs]

    @router.post("/material-catalogs")
    async def create_material_catalog(user=Depends(require_admin)):
        database = get_db()
        counter = await database.counters.find_one_and_update(
            {"_id": "material-catalogs"},
            {"$inc": {"seq": 1}},
            upsert=True,
            return_document=True,
        )
        seq = int((counter or {}).get("seq") or 1)
        now = datetime.now(timezone.utc).isoformat()
        doc = {
            "name": f"Catálogo {seq}",
            "materials": [],
            "created_at": now,
            "updated_at": now,
        }
        result = await database.material_catalogs.insert_one(doc)
        doc["_id"] = result.inserted_id
        await audit(user, "catalog.create", "material_catalog", str(result.inserted_id), {"name": doc["name"]})
        return serialize_catalog(doc)

    @router.delete("/material-catalogs/{catalog_id}")
    async def delete_material_catalog(catalog_id: str, user=Depends(require_admin)):
        if not ObjectId.is_valid(catalog_id):
            raise HTTPException(404, "Catálogo não encontrado")
        database = get_db()
        catalog = await database.material_catalogs.find_one({"_id": ObjectId(catalog_id)})
        if not catalog:
            raise HTTPException(404, "Catálogo não encontrado")
        linked = await database.categories.count_documents({"kit_catalog_ids": catalog_id})
        if linked:
            raise HTTPException(
                409,
                f"Este catálogo está vinculado a {linked} categoria(s). Remova o vínculo antes de excluir",
            )
        await database.material_catalogs.delete_one({"_id": catalog["_id"]})
        await audit(user, "catalog.delete", "material_catalog", catalog_id, {"name": catalog.get("name", "")})
        return {"message": "Catálogo removido"}

    @router.get("/material-catalogs/template")
    async def catalog_template(format: Literal["xlsx", "csv"] = "xlsx", user=Depends(require_admin)):
        examples = [["HGU5 CV", "Unidade", 1], ["Cabo Drop", "Metro", 500]]
        if format == "csv":
            buf = StringIO()
            writer = csv.writer(buf, delimiter=";")
            writer.writerows([CATALOG_HEADERS, *examples])
            content, media = buf.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Itens do catálogo"
            for row in [CATALOG_HEADERS, *examples]:
                ws.append(row)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="660099")
                ws.column_dimensions[cell.column_letter].width = 38
            stream = BytesIO()
            wb.save(stream)
            content, media = stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return Response(
            content,
            media_type=media,
            headers={"Content-Disposition": f'attachment; filename="modelo-catalogo.{format}"'},
        )

    @router.post("/material-catalogs/{catalog_id}/import")
    async def import_into_material_catalog(
        catalog_id: str,
        file: UploadFile = File(...),
        preview: bool = Form(True),
        user=Depends(require_admin),
    ):
        if not ObjectId.is_valid(catalog_id):
            raise HTTPException(404, "Catálogo não encontrado")
        raw = await file.read(5 * 1024 * 1024 + 1)
        if not raw or len(raw) > 5 * 1024 * 1024:
            raise HTTPException(400, "Envie uma planilha não vazia de até 5 MB")
        try:
            rows, errors = await asyncio.to_thread(parse_catalog_spreadsheet, raw, file.filename or "")
        except Exception as exc:
            raise HTTPException(
                400,
                str(exc) if isinstance(exc, ValueError) else "Planilha inválida. Salve novamente como .xlsx ou .csv",
            )
        database = get_db()
        catalog = await database.material_catalogs.find_one({"_id": ObjectId(catalog_id)})
        if not catalog:
            raise HTTPException(404, "Catálogo não encontrado")
        items, entries, counts, conflicts = plan_catalog_items(rows, catalog)
        errors.extend(conflicts)
        result = {
            "preview": preview,
            "valid": not errors,
            "rows": len(rows),
            "entries": entries,
            "errors": errors,
            **counts,
        }
        if preview or errors:
            return result
        if items != catalog.get("materials", []):
            await save_catalog_materials(database, catalog, items)
        await audit(user, "catalog.import", "material_catalog", catalog_id, counts)
        return result

    @router.post("/material-catalogs/{catalog_id}/materials")
    async def create_catalog_material(catalog_id: str, payload: MaterialInput, user=Depends(require_admin)):
        return await save_catalog_item(catalog_id, None, payload, user)

    @router.put("/material-catalogs/{catalog_id}/materials/{item_id}")
    async def update_catalog_material(catalog_id: str, item_id: str, payload: MaterialInput, user=Depends(require_admin)):
        return await save_catalog_item(catalog_id, item_id, payload, user)

    @router.post("/material-catalogs/{catalog_id}/materials/batch")
    async def create_catalog_materials_batch(catalog_id: str, payload: MaterialBatchInput, user=Depends(require_admin)):
        if not ObjectId.is_valid(catalog_id):
            raise HTTPException(404, "Catálogo não encontrado")
        database = get_db()
        catalog = await database.material_catalogs.find_one({"_id": ObjectId(catalog_id)})
        if not catalog:
            raise HTTPException(404, "Catálogo não encontrado")
        items = list(catalog.get("materials", []))
        if len(items) + len(payload.items) > MAX_ROWS:
            raise HTTPException(400, f"Limite de {MAX_ROWS} sub-itens por catálogo")

        existing_names = {name_key(item["name"]) for item in items}
        incoming_names = set()
        created = []
        for index, material in enumerate(payload.items, 1):
            data = material.model_dump()
            key = name_key(data["name"])
            if key in existing_names:
                raise HTTPException(409, f"Linha {index}: {data['name']} já está cadastrado neste catálogo")
            if key in incoming_names:
                raise HTTPException(409, f"Linha {index}: {data['name']} está repetido na montagem")
            incoming_names.add(key)
            created.append({"id": str(uuid.uuid4()), **data})

        await save_catalog_materials(database, catalog, [*items, *created])
        await audit(
            user,
            "catalog.batch_create",
            "material_catalog",
            catalog_id,
            {"count": len(created), "item_ids": [item["id"] for item in created]},
        )
        return {"created": len(created), "items": created}

    async def save_catalog_item(catalog_id, item_id, payload, user):
        if not ObjectId.is_valid(catalog_id):
            raise HTTPException(404, "Catálogo não encontrado")
        database = get_db()
        catalog = await database.material_catalogs.find_one({"_id": ObjectId(catalog_id)})
        if not catalog:
            raise HTTPException(404, "Catálogo não encontrado")
        items = list(catalog.get("materials", []))
        if item_id and not any(item["id"] == item_id for item in items):
            raise HTTPException(404, "Sub-item não encontrado")
        if any(name_key(item["name"]) == name_key(payload.name) and item["id"] != item_id for item in items):
            raise HTTPException(409, "Já existe este sub-item no catálogo")
        item = {"id": item_id or str(uuid.uuid4()), **payload.model_dump()}
        if item_id:
            items = [item if old["id"] == item_id else old for old in items]
        else:
            if len(items) >= MAX_ROWS:
                raise HTTPException(400, f"Limite de {MAX_ROWS} sub-itens por catálogo")
            items.append(item)
        await save_catalog_materials(database, catalog, items)
        await audit(
            user,
            "catalog.item_update" if item_id else "catalog.item_create",
            "material_catalog",
            catalog_id,
            {"item_id": item["id"], "name": item["name"]},
        )
        return item

    @router.get("/materials/template")
    async def template(format: Literal["xlsx", "csv"] = "xlsx", user=Depends(require_admin)):
        examples = [["Drop", "Drop Externo", "Metro", 500], ["Drop", "Drop Interno", "Metro", 100], ["HGU", "HGU Wi-Fi", "Unidade", 1]]
        if format == "csv":
            buf = StringIO()
            writer = csv.writer(buf, delimiter=";")
            writer.writerows([HEADERS, *examples])
            content, media = buf.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "Materiais"
            for row in [HEADERS, *examples]:
                ws.append(row)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="660099")
                ws.column_dimensions[cell.column_letter].width = 36
            stream = BytesIO()
            wb.save(stream)
            content, media = stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return Response(content, media_type=media, headers={"Content-Disposition": f'attachment; filename="modelo-materiais.{format}"'})

    @router.post("/materials/import")
    async def import_catalog(file: UploadFile = File(...), preview: bool = Form(True), user=Depends(require_admin)):
        raw = await file.read(5 * 1024 * 1024 + 1)
        if not raw or len(raw) > 5 * 1024 * 1024:
            raise HTTPException(400, "Envie uma planilha não vazia de até 5 MB")
        try:
            rows, errors = await asyncio.to_thread(parse_spreadsheet, raw, file.filename or "")
        except Exception as exc:
            raise HTTPException(400, str(exc) if isinstance(exc, ValueError) else "Planilha inválida. Salve novamente como .xlsx ou .csv")
        database = get_db()
        categories = await database.categories.find({}).to_list(10000)
        plans, entries, counts, conflicts = catalog_plan(rows, categories)
        errors.extend(conflicts)
        if len(categories) + counts["categories_created"] > 200:
            errors.append({"row": 1, "message": "Limite de 200 categorias. Agrupe os sub-itens em categorias existentes"})
        result = {"preview": preview, "valid": not errors, "rows": len(rows), "entries": entries, "errors": errors, **counts}
        if preview or errors:
            return result
        # Validation completes before the first write. Existing rules are never overwritten.
        try:
            for plan in plans:
                category = plan["category"]
                if plan["new"]:
                    await database.categories.insert_one({**category, "materials": plan["materials"]})
                elif plan["materials"] != category.get("materials", []):
                    await save_materials(database, category, plan["materials"])
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(409, "Importação interrompida. Atualize e reenvie a planilha; itens já salvos serão reconhecidos sem duplicação")
        await audit(user, "materials.import", "catalog", "materials", counts)
        return result

    @router.post("/categories/{category_id}/materials")
    async def create_material(category_id: str, payload: MaterialInput, user=Depends(require_admin)):
        return await save_one(category_id, None, payload, user)

    @router.put("/categories/{category_id}/materials/{item_id}")
    async def update_material(category_id: str, item_id: str, payload: MaterialInput, user=Depends(require_admin)):
        return await save_one(category_id, item_id, payload, user)

    @router.post("/categories/{category_id}/materials/batch")
    async def create_materials_batch(category_id: str, payload: MaterialBatchInput, user=Depends(require_admin)):
        if not ObjectId.is_valid(category_id):
            raise HTTPException(404, "Categoria não encontrada")
        database = get_db()
        category = await database.categories.find_one({"_id": ObjectId(category_id)})
        if not category:
            raise HTTPException(404, "Categoria não encontrada")
        items = list(category.get("materials", []))
        if len(items) + len(payload.items) > MAX_ROWS:
            raise HTTPException(400, f"Limite de {MAX_ROWS} sub-itens por categoria")

        existing_names = {name_key(item["name"]) for item in items}
        incoming_names = set()
        created = []
        for index, material in enumerate(payload.items, 1):
            data = material.model_dump()
            key = name_key(data["name"])
            if key in existing_names:
                raise HTTPException(409, f"Linha {index}: {data['name']} já está cadastrado nesta categoria")
            if key in incoming_names:
                raise HTTPException(409, f"Linha {index}: {data['name']} está repetido na montagem")
            incoming_names.add(key)
            created.append({"id": str(uuid.uuid4()), **data})

        await save_materials(database, category, [*items, *created])
        await audit(user, "material.batch_create", "category", category_id, {
            "count": len(created),
            "item_ids": [item["id"] for item in created],
        })
        return {"created": len(created), "items": created}

    async def save_one(category_id, item_id, payload, user):
        if not ObjectId.is_valid(category_id):
            raise HTTPException(404, "Categoria não encontrada")
        database = get_db()
        category = await database.categories.find_one({"_id": ObjectId(category_id)})
        if not category:
            raise HTTPException(404, "Categoria não encontrada")
        items = list(category.get("materials", []))
        if item_id and not any(item["id"] == item_id for item in items):
            raise HTTPException(404, "Sub-item não encontrado")
        if any(name_key(item["name"]) == name_key(payload.name) and item["id"] != item_id for item in items):
            raise HTTPException(409, "Já existe este sub-item na categoria")
        item = {"id": item_id or str(uuid.uuid4()), **payload.model_dump()}
        if item_id:
            items = [item if old["id"] == item_id else old for old in items]
        else:
            if len(items) >= MAX_ROWS:
                raise HTTPException(400, f"Limite de {MAX_ROWS} sub-itens por categoria")
            items.append(item)
        await save_materials(database, category, items)
        await audit(user, "material.update" if item_id else "material.create", "category", category_id, {"item_id": item["id"], "name": item["name"]})
        return item



async def validate_catalog_kit(database, entries):
    """Validate quantities against standalone material catalogs."""
    if not isinstance(entries, list) or not 1 <= len(entries) <= 500:
        raise HTTPException(400, "Selecione de 1 a 500 sub-itens para o kit")
    ids = {str(entry.get("category_id", "")) for entry in entries if isinstance(entry, dict)}
    if any(not ObjectId.is_valid(value) for value in ids):
        raise HTTPException(400, "Catálogo inválido no kit")
    catalogs = await database.material_catalogs.find(
        {"_id": {"$in": [ObjectId(value) for value in ids]}}
    ).to_list(200)
    by_id = {str(catalog["_id"]): catalog for catalog in catalogs}
    snapshots, seen = [], set()

    for entry in entries:
        if not isinstance(entry, dict):
            raise HTTPException(400, "Sub-item inválido no kit")
        catalog = by_id.get(str(entry.get("category_id", "")))
        item = next(
            (
                material
                for material in (catalog or {}).get("materials", [])
                if material["id"] == entry.get("item_id") and material.get("active", True)
            ),
            None,
        )
        if not item:
            raise HTTPException(400, "Um dos sub-itens está indisponível. Atualize o catálogo")
        key = (str(catalog["_id"]), item["id"])
        if key in seen:
            raise HTTPException(400, "O mesmo sub-item não pode aparecer duas vezes no kit")
        seen.add(key)
        quantity = entry.get("quantity")
        step = item["multiple"] if item["measure"] == "metro" else 1
        if type(quantity) is not int or not 1 <= quantity <= MAX_QUANTITY or quantity % step:
            raise HTTPException(400, f"{item['name']}: informe um inteiro positivo em múltiplos de {step}")
        snapshots.append({
            "category_id": str(catalog["_id"]),
            "category_name": catalog["name"],
            "item_id": item["id"],
            "name": item["name"],
            "measure": item["measure"],
            "multiple": step,
            "quantity": quantity,
        })
    return snapshots, catalogs

async def validate_kit(database, entries):
    if not isinstance(entries, list) or not 1 <= len(entries) <= 500:
        raise HTTPException(400, "Selecione de 1 a 500 sub-itens para o kit")
    ids = {str(entry.get("category_id", "")) for entry in entries if isinstance(entry, dict)}
    if any(not ObjectId.is_valid(value) for value in ids):
        raise HTTPException(400, "Categoria inválida no kit")
    categories = await database.categories.find({"_id": {"$in": [ObjectId(value) for value in ids]}, "active": True}).to_list(200)
    by_id = {str(category["_id"]): category for category in categories}
    snapshots, seen = [], set()
    for entry in entries:
        if not isinstance(entry, dict):
            raise HTTPException(400, "Sub-item inválido no kit")
        cat = by_id.get(str(entry.get("category_id", "")))
        item = next((i for i in (cat or {}).get("materials", []) if i["id"] == entry.get("item_id") and i.get("active", True)), None)
        if not item:
            raise HTTPException(400, "Um dos sub-itens está indisponível. Atualize o catálogo")
        key = (str(cat["_id"]), item["id"])
        if key in seen:
            raise HTTPException(400, "O mesmo sub-item não pode aparecer duas vezes no kit")
        seen.add(key)
        quantity = entry.get("quantity")
        step = item["multiple"] if item["measure"] == "metro" else 1
        if type(quantity) is not int or not 1 <= quantity <= MAX_QUANTITY or quantity % step:
            raise HTTPException(400, f"{item['name']}: informe um inteiro positivo em múltiplos de {step}")
        snapshots.append({"category_id": str(cat["_id"]), "category_name": cat["name"], "item_id": item["id"], "name": item["name"], "measure": item["measure"], "multiple": step, "quantity": quantity})
    return snapshots, categories
