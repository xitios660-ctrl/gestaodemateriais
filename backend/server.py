from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import logging
import asyncio
import csv
import uuid
import json
import random
import secrets
import hashlib
import ipaddress
import re
import smtplib
import ssl
from email.message import EmailMessage
from email.utils import formataddr
from datetime import datetime, timezone, timedelta
from typing import List, Optional, Annotated, Literal
from html import escape
from html.parser import HTMLParser
from urllib.parse import urlparse
from io import BytesIO, StringIO

import bcrypt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import jwt
import httpx
from fastapi import FastAPI, APIRouter, Request, Response, HTTPException, Depends, UploadFile, File, Form, BackgroundTasks
from starlette.middleware.cors import CORSMiddleware
from starlette.middleware.gzip import GZipMiddleware
from starlette.responses import Response as StarletteResponse, FileResponse, JSONResponse
from pydantic import BaseModel, Field, BeforeValidator, ConfigDict, EmailStr
from bson import ObjectId
from database import create_database
from materials import register_material_routes, validate_kit

# ----------------------------------------------------------------------------
# Setup
# ----------------------------------------------------------------------------
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

client, db, db_name, DB_ENGINE = create_database()

app = FastAPI(title="Gestão de Materiais API", version="1.0.0")
api_router = APIRouter(prefix="/api")

JWT_ALGORITHM = "HS256"
FRONTEND_URL = (os.environ.get("FRONTEND_URL") or os.environ.get("RENDER_EXTERNAL_URL") or "http://localhost:3000").rstrip("/")
LOCAL_STORAGE_DIR = Path(os.environ.get("STORAGE_DIR", ROOT_DIR / "uploads"))
LOCAL_STORAGE_DIR.mkdir(parents=True, exist_ok=True)

# Object storage
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_NAME = "central-servicos"
storage_key = None

# Email
EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ.get("EMERGENT_EMAIL_KEY", "")
EMAIL_FROM_NAME = os.environ.get("EMAIL_FROM_NAME") or "Central de Serviços"
SMTP_HOST = (os.environ.get("SMTP_HOST") or "").strip()
SMTP_PORT = int(os.environ.get("SMTP_PORT") or "587")
SMTP_USERNAME = (os.environ.get("SMTP_USERNAME") or "").strip()
SMTP_PASSWORD = os.environ.get("SMTP_PASSWORD") or ""
SMTP_FROM_EMAIL = (os.environ.get("SMTP_FROM_EMAIL") or SMTP_USERNAME).strip()
SMTP_STARTTLS = (os.environ.get("SMTP_STARTTLS") or "true").strip().lower() in ("1", "true", "yes")
BREVO_API_KEY = (os.environ.get("BREVO_API_KEY") or "").strip()
BREVO_SENDER_EMAIL = (os.environ.get("BREVO_SENDER_EMAIL") or "").strip().lower()
BREVO_SENDER_NAME = (os.environ.get("BREVO_SENDER_NAME") or EMAIL_FROM_NAME).strip()
BREVO_WEBHOOK_SECRET = (os.environ.get("BREVO_WEBHOOK_SECRET") or "").strip()
DEFAULT_OWNER_EMAIL = (os.environ.get("DEFAULT_OWNER_EMAIL") or "").strip().lower()
PRIMARY_NOTIFICATION_EMAIL = (os.environ.get("PRIMARY_NOTIFICATION_EMAIL") or DEFAULT_OWNER_EMAIL).strip().lower()
DEFAULT_OWNERS = [DEFAULT_OWNER_EMAIL] if DEFAULT_OWNER_EMAIL else []

STATUS_LABELS = {
    "aberto": "Aberto",
    "em_analise": "Em Análise",
    "em_andamento": "Em Andamento",
    "concluido": "Concluído",
    "cancelado": "Cancelado",
}

# ----------------------------------------------------------------------------
# Database helpers
# ----------------------------------------------------------------------------
def _validate_object_id(v):
    if isinstance(v, ObjectId):
        return str(v)
    return str(v)

PyObjectId = Annotated[str, BeforeValidator(_validate_object_id)]

MIN_PASSWORD_LENGTH = 8


def as_object_id(value: str, detail: str = "Recurso não encontrado") -> ObjectId:
    try:
        return ObjectId(str(value))
    except Exception:
        raise HTTPException(status_code=404, detail=detail)


def get_client_ip(request: Request) -> str:
    """Return a stable client IP behind Render/Cloudflare proxies."""
    candidates = [
        request.headers.get("cf-connecting-ip"),
        (request.headers.get("x-forwarded-for") or "").split(",")[0].strip(),
        request.client.host if request.client else "",
    ]
    for candidate in candidates:
        if not candidate:
            continue
        try:
            return str(ipaddress.ip_address(candidate.strip()))
        except ValueError:
            continue
    return "unknown"


async def enforce_rate_limit(request: Request, scope: str, limit: int, minutes: int = 1) -> None:
    """Small Mongo-backed limiter for unauthenticated, abuse-prone endpoints."""
    ip = get_client_ip(request)
    now = datetime.now(timezone.utc)
    bucket = int(now.timestamp() // max(60, minutes * 60))
    key = hashlib.sha256(f"{scope}:{ip}:{bucket}".encode()).hexdigest()
    doc = await db.rate_limits.find_one_and_update(
        {"_id": key},
        {
            "$inc": {"count": 1},
            "$setOnInsert": {
                "scope": scope,
                "expires_at": now + timedelta(minutes=max(2, minutes * 2)),
            },
        },
        upsert=True,
        return_document=True,
    )
    if doc and int(doc.get("count", 0)) > limit:
        raise HTTPException(status_code=429, detail="Muitas solicitações. Aguarde um pouco e tente novamente.")


class BaseDocument(BaseModel):
    model_config = ConfigDict(populate_by_name=True, arbitrary_types_allowed=True)
    id: Optional[PyObjectId] = Field(default=None, alias="_id")

    @classmethod
    def from_mongo(cls, doc: dict):
        if not doc:
            return None
        return cls(**doc)


# ----------------------------------------------------------------------------
# Password / JWT helpers
# ----------------------------------------------------------------------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False


def get_jwt_secret() -> str:
    secret = os.environ.get("JWT_SECRET")
    if secret:
        return secret
    # Development fallback only. Production deployments should always set JWT_SECRET.
    if os.environ.get("RENDER") or os.environ.get("ENVIRONMENT") == "production":
        raise RuntimeError("JWT_SECRET não configurado")
    return "dev-only-change-me-please-32-characters"


def create_access_token(user_id: str, email: str, token_version: int = 0) -> str:
    payload = {"sub": user_id, "email": email, "ver": token_version,
               "exp": datetime.now(timezone.utc) + timedelta(minutes=15), "type": "access"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)


def create_refresh_token(user_id: str, token_version: int = 0) -> str:
    payload = {"sub": user_id, "ver": token_version,
               "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "refresh"}
    return jwt.encode(payload, get_jwt_secret(), algorithm=JWT_ALGORITHM)


def set_auth_cookies(response: Response, access: str, refresh: str):
    force_secure = os.environ.get("COOKIE_SECURE")
    if force_secure is None:
        secure = FRONTEND_URL.startswith("https://")
    else:
        secure = force_secure.lower() in ("1", "true", "yes")
    same_site = "lax"
    response.set_cookie(key="access_token", value=access, httponly=True, secure=secure,
                        samesite=same_site, max_age=900, path="/")
    response.set_cookie(key="refresh_token", value=refresh, httponly=True, secure=secure,
                        samesite=same_site, max_age=604800, path="/")


async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Não autenticado")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "access":
            raise HTTPException(status_code=401, detail="Tipo de token inválido")
        user = await db.users.find_one({"_id": as_object_id(payload["sub"], "Sessão inválida")})
        if not user:
            raise HTTPException(status_code=401, detail="Usuário não encontrado")
        if payload.get("ver", 0) != user.get("token_version", 0):
            raise HTTPException(status_code=401, detail="Sessão expirada")
        user["_id"] = str(user["_id"])
        user.pop("password_hash", None)
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expirado")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token inválido")


async def require_admin(user: dict = Depends(get_current_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Acesso restrito a administradores")
    return user


async def validate_category_ids(category_ids: List[str]) -> List[str]:
    normalized = list(dict.fromkeys(str(value) for value in (category_ids or [])))
    if not normalized:
        return []
    if any(not ObjectId.is_valid(value) for value in normalized):
        raise HTTPException(status_code=400, detail="Uma das categorias informadas é inválida")
    object_ids = [ObjectId(value) for value in normalized]
    found = await db.categories.count_documents({"_id": {"$in": object_ids}})
    if found != len(object_ids):
        raise HTTPException(status_code=400, detail="Uma das categorias informadas não existe")
    return normalized


async def validate_kit_catalog_ids(catalog_ids: List[str]) -> List[str]:
    """Validate links using the existing categories collection; no migration/new collection."""
    normalized = await validate_category_ids(catalog_ids)
    if not normalized:
        return []
    object_ids = [ObjectId(value) for value in normalized]
    docs = await db.categories.find({"_id": {"$in": object_ids}}).to_list(50)
    usable = {
        str(doc["_id"])
        for doc in docs
        if doc.get("active", True)
        and any(item.get("active", True) for item in (doc.get("materials") or []))
    }
    unavailable = [value for value in normalized if value not in usable]
    if unavailable:
        raise HTTPException(
            status_code=400,
            detail="Selecione somente catálogos que possuam ao menos um sub-item disponível",
        )
    return normalized


def scope_match(user: dict) -> dict:
    """Empty for admins; category filter for scoped responsáveis."""
    if user.get("role") == "admin":
        return {}
    return ticket_category_match(user.get("categories", []) or [])


def ticket_category_match(ids):
    return {"$or": [{"category_id": {"$in": ids}}, {"category_ids": {"$in": ids}}]}


def can_access_ticket(user, ticket):
    return user.get("role") == "admin" or bool(
        set(user.get("categories") or []).intersection(
            ticket.get("category_ids") or [ticket.get("category_id")]
        )
    )


async def get_all_admin_notification_emails() -> List[str]:
    admins = await db.users.find(
        {"role": "admin"},
        {"email": 1},
    ).to_list(500)
    return unique_recipient_emails([
        str(admin.get("email") or "").strip().lower()
        for admin in admins
        if admin.get("email")
    ])


async def get_ticket_notification_admin_email() -> str:
    """Return the admin account that most recently logged in to manage the panel."""
    target = await db.system_settings.find_one({"_id": "ticket_notification_admin"})
    if target:
        user_id = str(target.get("user_id") or "")
        if ObjectId.is_valid(user_id):
            admin = await db.users.find_one({"_id": ObjectId(user_id), "role": "admin"})
            if admin and admin.get("email"):
                return str(admin["email"]).strip().lower()

    # Backward-compatible fallback until an administrator logs in after this update.
    if PRIMARY_NOTIFICATION_EMAIL:
        return PRIMARY_NOTIFICATION_EMAIL
    admin = await db.users.find_one({"role": "admin"}, sort=[("last_login_at", -1), ("created_at", 1)])
    return str((admin or {}).get("email") or "").strip().lower()


async def write_audit(user: dict, action: str, entity_type: str, entity_id: str, details: Optional[dict] = None):
    """Best-effort audit trail. Never store credentials or raw secrets here."""
    try:
        await db.audit_logs.insert_one({
            "actor_id": str(user.get("_id", "")),
            "actor_email": user.get("email", ""),
            "actor_name": user.get("name", ""),
            "action": action,
            "entity_type": entity_type,
            "entity_id": str(entity_id),
            "details": details or {},
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    except Exception as exc:
        logger.warning("Falha ao registrar auditoria (%s): %s", action, exc)


# ----------------------------------------------------------------------------
# Object storage helpers
# ----------------------------------------------------------------------------
def _safe_local_path(path: str) -> Path:
    clean = path.lstrip("/").replace("\\", "/")
    target = (LOCAL_STORAGE_DIR / clean).resolve()
    root = LOCAL_STORAGE_DIR.resolve()
    if root not in target.parents and target != root:
        raise ValueError("Caminho de arquivo inválido")
    return target


async def init_storage(force: bool = False):
    """Use Emergent object storage when configured; otherwise use local disk."""
    global storage_key
    if not EMERGENT_KEY:
        return None
    if storage_key and not force:
        return storage_key
    async with httpx.AsyncClient(timeout=30) as http:
        resp = await http.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY})
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


async def put_object(path: str, data: bytes, content_type: str) -> dict:
    try:
        if not EMERGENT_KEY:
            target = _safe_local_path(path)
            target.parent.mkdir(parents=True, exist_ok=True)
            await asyncio.to_thread(target.write_bytes, data)
            return {"path": path, "storage": "local", "size": len(data)}

        key = await init_storage()
        async with httpx.AsyncClient(timeout=120) as http:
            resp = await http.put(
                f"{STORAGE_URL}/objects/{path}",
                headers={"X-Storage-Key": key, "Content-Type": content_type},
                content=data,
            )
            if resp.status_code == 404:
                key = await init_storage(force=True)
                resp = await http.put(
                    f"{STORAGE_URL}/objects/{path}",
                    headers={"X-Storage-Key": key, "Content-Type": content_type},
                    content=data,
                )
        resp.raise_for_status()
        return resp.json()
    except HTTPException:
        raise
    except (OSError, httpx.HTTPError, KeyError, ValueError) as exc:
        logger.error("Falha ao armazenar arquivo: %s", exc)
        raise HTTPException(status_code=503, detail="Não foi possível armazenar o arquivo agora. Tente novamente.")


async def get_object(path: str):
    try:
        if not EMERGENT_KEY:
            target = _safe_local_path(path)
            if not target.exists() or not target.is_file():
                raise HTTPException(status_code=404, detail="Arquivo não encontrado")
            data = await asyncio.to_thread(target.read_bytes)
            return data, "application/octet-stream"

        key = await init_storage()
        async with httpx.AsyncClient(timeout=60) as http:
            resp = await http.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key})
            if resp.status_code == 404:
                key = await init_storage(force=True)
                resp = await http.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key})
        if resp.status_code == 404:
            raise HTTPException(status_code=404, detail="Arquivo não encontrado")
        resp.raise_for_status()
        return resp.content, resp.headers.get("Content-Type", "application/octet-stream")
    except HTTPException:
        raise
    except (OSError, httpx.HTTPError, KeyError, ValueError) as exc:
        logger.error("Falha ao recuperar arquivo: %s", exc)
        raise HTTPException(status_code=503, detail="Não foi possível acessar o arquivo agora. Tente novamente.")


# ----------------------------------------------------------------------------
# Email helpers (guardrail gate copied from playbook)
# ----------------------------------------------------------------------------
_SHORTENERS = ("bit.ly", "tinyurl.com", "t.co", "is.gd", "cutt.ly", "goo.gl", "rebrand.ly")
_CRED_ASK = ("reply with your password", "reply with the code", "send your password", "cvv",
             "send us your password", "enter your password below", "confirm your card number",
             "your full card number", "seed phrase", "recovery phrase", "verify your card",
             "social security number", "confirm your bank details")
_HOSTISH = re.compile(r"\b(?:https?://)?((?:[a-z0-9-]+\.)+[a-z]{2,})", re.I)


def _host_ok(host: str) -> bool:
    if not host or "xn--" in host:
        return False
    try:
        ipaddress.ip_address(host)
        return False
    except ValueError:
        pass
    return not any(host == s or host.endswith("." + s) for s in _SHORTENERS)


def _same_site(shown: str, real: str) -> bool:
    return shown == real or real.endswith("." + shown) or shown.endswith("." + real)


class _EmailScan(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.urls, self.anchors = set(), [], []
        self._href, self._text = None, []

    def handle_starttag(self, tag, attrs):
        self.tags.add(tag.lower())
        self.urls += [v for k, v in attrs if k.lower() in ("href", "src") and v]
        if tag.lower() == "a":
            self._href = dict((k.lower(), v) for k, v in attrs).get("href")
            self._text = []

    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag):
        if tag.lower() == "a" and self._href is not None:
            self.anchors.append((self._href, "".join(self._text)))
            self._href, self._text = None, []


def _assert_safe_email(subject: str, html: str) -> None:
    scan = _EmailScan()
    scan.feed(html)
    if scan.tags & {"form", "input", "textarea", "select"}:
        raise ValueError("No forms or input fields in email (G2)")
    body = f"{subject}\n{html}".lower()
    for p in _CRED_ASK:
        if p in body:
            raise ValueError(f"Email asks for credentials: {p!r} (G2)")
    for url in scan.urls:
        low = url.strip().lower()
        if low.startswith(("mailto:", "tel:", "cid:", "#")):
            continue
        if not low.startswith("https://"):
            raise ValueError(f"Email links must be absolute https: {url!r} (G3)")
        host = urlparse(low).hostname or ""
        if not _host_ok(host) or urlparse(low).username is not None:
            raise ValueError(f"Bad URL: {url!r} (G3)")
    for href, text in scan.anchors:
        real = urlparse(href.strip().lower()).hostname or ""
        if not real:
            continue
        for m in _HOSTISH.finditer(text):
            if not _same_site(m.group(1).lower(), real):
                raise ValueError(f"Anchor text mismatch (G3)")


def _email_identity(value: str) -> str:
    email = str(value or "").strip().lower()
    if "@" not in email:
        return email
    local, domain = email.rsplit("@", 1)
    if domain in {"gmail.com", "googlemail.com"}:
        local = local.split("+", 1)[0].replace(".", "")
        domain = "gmail.com"
    return f"{local}@{domain}"


def unique_recipient_emails(values) -> List[str]:
    result, seen = [], set()
    for value in values:
        email = str(value or "").strip().lower()
        if not email:
            continue
        identity = _email_identity(email)
        if identity in seen:
            continue
        seen.add(identity)
        result.append(email)
    return result


def email_delivery_configured() -> bool:
    brevo_ready = bool(BREVO_API_KEY and BREVO_SENDER_EMAIL)
    emergent_ready = bool(EMAIL_KEY and not EMAIL_KEY.startswith("{"))
    smtp_ready = bool(SMTP_HOST and SMTP_FROM_EMAIL and SMTP_USERNAME and SMTP_PASSWORD)
    return brevo_ready or emergent_ready or smtp_ready


def _smtp_send(to: str, subject: str, html: str) -> str:
    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = formataddr((EMAIL_FROM_NAME, SMTP_FROM_EMAIL))
    message["To"] = to
    message.set_content("Este e-mail contém conteúdo em HTML. Abra-o em um cliente compatível.")
    message.add_alternative(html, subtype="html")

    context = ssl.create_default_context()
    if SMTP_PORT == 465:
        with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT, timeout=25, context=context) as smtp:
            smtp.login(SMTP_USERNAME, SMTP_PASSWORD.replace(" ", ""))
            smtp.send_message(message)
    else:
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=25) as smtp:
            smtp.ehlo()
            if SMTP_STARTTLS:
                smtp.starttls(context=context)
                smtp.ehlo()
            smtp.login(SMTP_USERNAME, SMTP_PASSWORD.replace(" ", ""))
            smtp.send_message(message)
    return f"smtp:{uuid.uuid4()}"


async def send_email(to: str, subject: str, html: str) -> Optional[str]:
    _assert_safe_email(subject, html)

    if BREVO_API_KEY and BREVO_SENDER_EMAIL:
        try:
            payload = {
                "sender": {"email": BREVO_SENDER_EMAIL, "name": BREVO_SENDER_NAME},
                "to": [{"email": to}],
                "subject": subject,
                "htmlContent": html,
            }
            async with httpx.AsyncClient(timeout=30) as client_http:
                resp = await client_http.post(
                    "https://api.brevo.com/v3/smtp/email",
                    headers={
                        "api-key": BREVO_API_KEY,
                        "accept": "application/json",
                        "content-type": "application/json",
                    },
                    json=payload,
                )
            resp.raise_for_status()
            data = resp.json()
            return data.get("messageId") or data.get("messageIds", [None])[0]
        except Exception as exc:
            logger.error("Brevo send error to %s: %s: %s", to, type(exc).__name__, str(exc)[:240])

    if EMAIL_KEY and not EMAIL_KEY.startswith("{"):
        payload = {"to": [to], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
        try:
            async with httpx.AsyncClient(timeout=30) as client_http:
                resp = await client_http.post(
                    f"{EMAIL_BASE_URL}/api/v1/email/send",
                    headers={"X-Email-Key": EMAIL_KEY},
                    json=payload,
                )
            resp.raise_for_status()
            return resp.json().get("id")
        except Exception as exc:
            logger.warning("Provider de e-mail principal falhou; tentando SMTP: %s", type(exc).__name__)

    if SMTP_HOST and SMTP_FROM_EMAIL and SMTP_USERNAME and SMTP_PASSWORD:
        try:
            return await asyncio.to_thread(_smtp_send, to, subject, html)
        except Exception as exc:
            logger.error("SMTP send error to %s: %s: %s", to, type(exc).__name__, str(exc)[:240])
            return None

    logger.error("E-mail não configurado. Defina EMERGENT_EMAIL_KEY ou SMTP_HOST/SMTP_USERNAME/SMTP_PASSWORD/SMTP_FROM_EMAIL")
    return None


def format_field_value(value) -> str:
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, dict) and ("parent" in value or "child" in value or "children" in value):
        parent = str(value.get("parent") or "").strip()
        children = value.get("children")
        if not isinstance(children, list):
            legacy_child = str(value.get("child") or "").strip()
            children = [legacy_child] if legacy_child else []
        clean_children = [str(item).strip() for item in children if str(item).strip()]
        if parent and clean_children:
            return f"{parent} → {', '.join(clean_children)}"
        return parent or ", ".join(clean_children) or "—"
    return str(value if value not in (None, "") else "—")


async def notify_owners(ticket: dict, owners: List[str]):
    base = FRONTEND_URL.rstrip("/")
    if not base.startswith("https://"):
        base = "https://ticket-builder-14.preview.emergentagent.com"
    track_link = f"{base}/acompanhar?q={escape(ticket['ticket_number'])}"
    req = ticket.get("requester", {})
    rows = "".join(
        f'<tr><td style="padding:4px 12px 4px 0;color:#64748b">{escape(str(k))}</td>'
        f'<td style="padding:4px 0;color:#0f172a"><strong>{escape(str(v))}</strong></td></tr>'
        for k, v in [
            ("Chamado", ticket["ticket_number"]),
            ("Categoria", ticket.get("category_name", "")),
            ("Matrícula", req.get("matricula", "")),
            ("E-mail", req.get("email", "")),
            ("Empresa", req.get("empresa", "")),
            ("Prazo (SLA)", f"{ticket.get('lead_time_hours', 0)} horas"),
        ]
    )
    detail_rows = "".join(
        f'<tr><td style="padding:4px 12px 4px 0;color:#64748b">{escape(str(k))}</td>'
        f'<td style="padding:4px 0;color:#0f172a"><strong>{escape(format_field_value(v))}</strong></td></tr>'
        for k, v in (ticket.get("field_values") or {}).items()
    )
    rows += detail_rows
    rows += "".join(
        f'<tr><td style="padding:4px 12px 4px 0">{escape(i["category_name"])} / {escape(i["name"])}</td>'
        f'<td><strong>{i["quantity"]} {"m" if i["measure"] == "metro" else "un"}</strong></td></tr>'
        for i in ticket.get("material_items", [])
    )
    html = (
        f'<table role="presentation" width="100%"><tr><td style="padding:24px;font-family:Arial,sans-serif;max-width:560px">'
        f'<h2 style="color:#7c3aed;margin:0 0 16px">Novo chamado aberto: {escape(ticket["ticket_number"])}</h2>'
        f'<p style="color:#475569">Um novo chamado foi registrado na <strong>{escape(EMAIL_FROM_NAME)}</strong> e você está cadastrado como responsável.</p>'
        f'<table role="presentation" style="margin:16px 0;font-size:14px">{rows}</table>'
        f'<p><a href="{escape(track_link)}" style="background:#7c3aed;color:#fff;padding:10px 20px;'
        f'border-radius:8px;text-decoration:none;display:inline-block">Acompanhar chamado</a></p>'
        f'<p style="font-size:12px;color:#94a3b8;margin-top:24px">Enviado por {escape(EMAIL_FROM_NAME)}.</p>'
        f'</td></tr></table>'
    )
    log = []
    for owner in owners:
        eid = await send_email(owner, f"[{ticket['ticket_number']}] Novo chamado - {ticket.get('category_name','')}", html)
        log.append({"to": owner, "email_id": eid, "sent": bool(eid),
                    "at": datetime.now(timezone.utc).isoformat()})
    await db.tickets.update_one({"ticket_number": ticket["ticket_number"]},
                                {"$set": {"email_log": log}})


async def send_ticket_opened_email(ticket: dict) -> bool:
    requester = ticket.get("requester") or {}
    to_email = str(requester.get("email") or "").strip().lower()
    if not to_email:
        logger.warning("Opened-ticket e-mail skipped for %s: requester e-mail missing", ticket.get("ticket_number"))
        return False

    base = FRONTEND_URL.rstrip("/")
    if not base.startswith("https://"):
        logger.error("Opened-ticket e-mail skipped: FRONTEND_URL inválida")
        return False

    ticket_number = escape(str(ticket.get("ticket_number") or ""))
    category_name = escape(str(ticket.get("category_name") or ""))
    track_link = f"{base}/acompanhar?q={ticket_number}"
    html = (
        f'<table role="presentation" width="100%"><tr><td style="padding:24px;font-family:Arial,sans-serif;max-width:560px">'
        f'<h2 style="color:#660099;margin:0 0 16px">Chamado {ticket_number} aberto com sucesso</h2>'
        f'<p style="color:#475569">Recebemos seu chamado de <strong>{category_name}</strong>.</p>'
        f'<p style="color:#475569"><strong>Status atual:</strong> Aberto</p>'
        f'<p><a href="{escape(track_link)}" style="background:#660099;color:#fff;padding:10px 20px;'
        f'border-radius:8px;text-decoration:none;display:inline-block">Acompanhar chamado</a></p>'
        f'<p style="font-size:12px;color:#94a3b8;margin-top:24px">Enviado por {escape(EMAIL_FROM_NAME)}.</p>'
        f'</td></tr></table>'
    )
    email_id = await send_email(to_email, f"[{ticket.get('ticket_number','')}] Chamado aberto", html)
    await db.tickets.update_one(
        {"_id": ticket["_id"]},
        {"$push": {
            "status_email_log": {
                "to": to_email,
                "from_status": None,
                "to_status": "aberto",
                "email_id": email_id,
                "sent": bool(email_id),
                "at": datetime.now(timezone.utc).isoformat(),
            }
        }},
    )
    return bool(email_id)


async def send_status_update_email(ticket: dict, previous_status: str, new_status: str, note: Optional[str] = None, responder: Optional[dict] = None) -> bool:
    requester = ticket.get("requester") or {}
    to_email = str(requester.get("email") or "").strip().lower()
    if not to_email:
        logger.warning("Status e-mail skipped for %s: requester e-mail missing", ticket.get("ticket_number"))
        return False

    base = FRONTEND_URL.rstrip("/")
    if not base.startswith("https://"):
        logger.error("Status e-mail skipped: FRONTEND_URL inválida")
        return False

    ticket_number = escape(str(ticket.get("ticket_number") or ""))
    category_name = escape(str(ticket.get("category_name") or ""))
    old_label = escape(STATUS_LABELS.get(previous_status, previous_status or ""))
    new_label = escape(STATUS_LABELS.get(new_status, new_status or ""))
    track_link = f"{base}/acompanhar?q={ticket_number}"
    safe_note = escape(str(note or "").strip())
    note_html = (
        f'<p style="color:#475569"><strong>Observação:</strong> {safe_note}</p>'
        if safe_note else ""
    )
    responder_data = responder or ticket.get("assigned_to") or {}
    responder_name = escape(str(responder_data.get("name") or responder_data.get("email") or "Equipe de atendimento"))
    responder_email = escape(str(responder_data.get("email") or "").strip().lower())
    responder_email_html = (
        f'<div style="color:#64748b;font-size:13px;margin-top:2px">{responder_email}</div>'
        if responder_email else ""
    )
    responder_html = (
        f'<div style="margin:18px 0;padding:14px 16px;border-radius:10px;background:#f8f5fb">'
        f'<div style="font-size:12px;color:#7c3aed;text-transform:uppercase;font-weight:700">Quem está atendendo</div>'
        f'<div style="margin-top:5px;color:#1f2937;font-weight:700">{responder_name}</div>'
        f'{responder_email_html}'
        f'</div>'
    )

    html = (
        f'<table role="presentation" width="100%"><tr><td style="padding:24px;font-family:Arial,sans-serif;max-width:560px">'
        f'<h2 style="color:#660099;margin:0 0 16px">Atualização do chamado {ticket_number}</h2>'
        f'<p style="color:#475569">Seu chamado de <strong>{category_name}</strong> teve o status atualizado.</p>'
        f'<table role="presentation" style="margin:16px 0;font-size:14px">'
        f'<tr><td style="padding:4px 12px 4px 0;color:#64748b">Status anterior</td><td><strong>{old_label}</strong></td></tr>'
        f'<tr><td style="padding:4px 12px 4px 0;color:#64748b">Novo status</td><td><strong>{new_label}</strong></td></tr>'
        f'</table>'
        f'{responder_html}'
        f'{note_html}'
        f'<p><a href="{escape(track_link)}" style="background:#660099;color:#fff;padding:10px 20px;'
        f'border-radius:8px;text-decoration:none;display:inline-block">Acompanhar chamado</a></p>'
        f'<p style="font-size:12px;color:#94a3b8;margin-top:24px">Enviado por {escape(EMAIL_FROM_NAME)}.</p>'
        f'</td></tr></table>'
    )

    email_id = await send_email(
        to_email,
        f"[{ticket.get('ticket_number','')}] Status atualizado - {STATUS_LABELS.get(new_status, new_status)}",
        html,
    )
    await db.tickets.update_one(
        {"_id": ticket["_id"]},
        {"$push": {
            "status_email_log": {
                "to": to_email,
                "from_status": previous_status,
                "to_status": new_status,
                "email_id": email_id,
                "sent": bool(email_id),
                "at": datetime.now(timezone.utc).isoformat(),
            }
        }},
    )
    return bool(email_id)


async def send_password_reset_email(to_email: str, token: str) -> bool:
    base = FRONTEND_URL.rstrip("/")
    link = f"{base}/reset-password?token={token}"
    if not email_delivery_configured() or not base.startswith("https://"):
        if urlparse(base).hostname in ("localhost", "127.0.0.1", "::1"):
            logger.warning("Email not configured; reset link: %s", link)
        else:
            logger.error("Reset email not configured (EMERGENT_EMAIL_KEY / FRONTEND_URL)")
        return False
    brand = escape(EMAIL_FROM_NAME)
    html = (
        f'<table role="presentation" width="100%"><tr><td style="padding:24px;font-family:Arial,sans-serif;max-width:520px">'
        f'<h2 style="color:#660099;margin:0 0 16px">Redefinição de senha</h2>'
        f'<p style="color:#475569">Recebemos uma solicitação para redefinir a senha do painel <strong>{brand}</strong>.</p>'
        f'<p style="margin:20px 0"><a href="{escape(link)}" style="background:#660099;color:#fff;padding:12px 24px;'
        f'border-radius:8px;text-decoration:none;display:inline-block">Redefinir senha</a></p>'
        f'<p style="color:#475569">Este link expira em 1 hora e só pode ser usado uma vez. Se você não '
        f'solicitou, ignore este e-mail — sua senha permanece inalterada.</p>'
        f'<p style="font-size:12px;color:#94a3b8;margin-top:24px">Enviado por {brand}. Nunca pedimos sua senha por e-mail.</p>'
        f'</td></tr></table>'
    )
    try:
        eid = await send_email(to_email, f"Redefinição de senha - {EMAIL_FROM_NAME}", html)
        return bool(eid)
    except Exception as e:
        logger.error(f"Reset email failed: {e}")
        return False


async def send_welcome_email(to_email: str, name: str, token: str) -> bool:
    base = FRONTEND_URL.rstrip("/")
    link = f"{base}/reset-password?token={token}"
    if not email_delivery_configured() or not base.startswith("https://"):
        logger.warning("Welcome email not sent (config); set-password link: %s", link)
        return False
    brand = escape(EMAIL_FROM_NAME)
    html = (
        f'<table role="presentation" width="100%"><tr><td style="padding:24px;font-family:Arial,sans-serif;max-width:520px">'
        f'<h2 style="color:#660099;margin:0 0 16px">Bem-vindo(a) à {brand}</h2>'
        f'<p style="color:#475569">Olá {escape(name)}, foi criada uma conta de acesso ao painel <strong>{brand}</strong>.</p>'
        f'<p style="color:#475569">Seu login (e-mail de acesso): <strong>{escape(to_email)}</strong></p>'
        f'<p style="color:#475569">Para começar, defina a sua senha de acesso:</p>'
        f'<p style="margin:20px 0"><a href="{escape(link)}" style="background:#660099;color:#fff;padding:12px 24px;'
        f'border-radius:8px;text-decoration:none;display:inline-block">Definir minha senha</a></p>'
        f'<p style="color:#475569">Este link expira em 48 horas e só pode ser usado uma vez.</p>'
        f'<p style="font-size:12px;color:#94a3b8;margin-top:24px">Enviado por {brand}. Nunca pedimos sua senha por e-mail.</p>'
        f'</td></tr></table>'
    )
    try:
        eid = await send_email(to_email, f"Acesso ao painel {EMAIL_FROM_NAME}", html)
        return bool(eid)
    except Exception as e:
        logger.error(f"Welcome email failed: {e}")
        return False


# ----------------------------------------------------------------------------
# Pydantic models
# ----------------------------------------------------------------------------
class LoginInput(BaseModel):
    email: EmailStr
    password: str


class ForgotInput(BaseModel):
    email: EmailStr


class ResetInput(BaseModel):
    token: str
    password: str


class UserCreate(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    email: EmailStr
    password: Optional[str] = Field(default=None, max_length=256)
    role: Literal["admin", "responsavel"] = "responsavel"
    categories: List[str] = Field(default_factory=list, max_length=200)
    send_welcome: bool = True


class UserUpdate(BaseModel):
    name: Optional[str] = Field(default=None, max_length=120)
    role: Optional[Literal["admin", "responsavel"]] = None
    password: Optional[str] = Field(default=None, max_length=256)
    categories: Optional[List[str]] = Field(default=None, max_length=200)


class DependentOption(BaseModel):
    parent: str = Field(min_length=1, max_length=160)
    children: List[str] = Field(default_factory=list, min_length=1, max_length=200)


class CustomField(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()), max_length=80)
    label: str = Field(min_length=1, max_length=160)
    type: Literal["text", "textarea", "select", "number", "date", "checkbox", "dependent_select"]
    required: bool = False
    options: List[str] = Field(default_factory=list, max_length=100)
    parent_label: Optional[str] = Field(default=None, max_length=160)
    child_label: Optional[str] = Field(default=None, max_length=160)
    dependent_options: List[DependentOption] = Field(default_factory=list, max_length=200)


class CategoryInput(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    icon: str = Field(default="Laptop", max_length=80)
    description: str = Field(default="", max_length=1200)
    lead_time_hours: int = Field(default=24, ge=1, le=720)
    owners: List[EmailStr] = Field(default_factory=list, max_length=100)
    fields: List[CustomField] = Field(default_factory=list, max_length=100)
    template_columns: List[str] = Field(default_factory=list, max_length=100)
    template_filename: Optional[str] = Field(default=None, max_length=180)
    kit_enabled: bool = False
    kit_catalog_ids: List[str] = Field(default_factory=list, max_length=50)
    active: bool = True


class StatusUpdate(BaseModel):
    status: Literal["aberto", "em_analise", "em_andamento", "concluido", "cancelado"]
    note: Optional[str] = Field(default=None, max_length=1000)


# ----------------------------------------------------------------------------
# Serialization
# ----------------------------------------------------------------------------
def ser_category(doc: dict) -> dict:
    doc = dict(doc)
    doc["id"] = str(doc.pop("_id"))
    return doc


def ser_category_public(doc: dict) -> dict:
    """Public category shape. Internal routing e-mails must never leave the admin API."""
    data = ser_category(doc)
    data.pop("owners", None)
    data["materials"] = [item for item in data.get("materials", []) if item.get("active", True)]
    return data


def ser_ticket(doc: dict) -> dict:
    doc = dict(doc)
    doc["id"] = str(doc.pop("_id"))
    doc["status_label"] = STATUS_LABELS.get(doc.get("status"), doc.get("status"))
    return doc


def ser_ticket_public(doc: dict) -> dict:
    """Minimal ticket shape used by the unauthenticated tracking screen."""
    data = ser_ticket(doc)
    allowed = {
        "id", "ticket_number", "category_name", "category_icon", "status",
        "status_label", "lead_time_hours", "created_at", "due_at",
    }
    return {key: value for key, value in data.items() if key in allowed}


def ser_user(doc: dict) -> dict:
    return {
        "id": str(doc["_id"]),
        "name": doc.get("name", ""),
        "email": doc.get("email", ""),
        "role": doc.get("role", "admin"),
        "categories": doc.get("categories", []),
        "created_at": doc.get("created_at"),
    }


# ----------------------------------------------------------------------------
# Auth endpoints
# ----------------------------------------------------------------------------
LOCKOUT_ATTEMPTS = 5
LOCKOUT_MINUTES = 15


@api_router.post("/auth/login")
async def login(payload: LoginInput, request: Request, response: Response):
    email = payload.email.lower().strip()
    ip = get_client_ip(request)
    identifier = f"{ip}:{email}"
    window = datetime.now(timezone.utc) - timedelta(minutes=LOCKOUT_MINUTES)
    recent = await db.login_attempts.count_documents(
        {"identifier": identifier, "at": {"$gt": window.isoformat()}})
    if recent >= LOCKOUT_ATTEMPTS:
        raise HTTPException(status_code=429, detail="Muitas tentativas. Tente novamente em 15 minutos.")

    user = await db.users.find_one({"email": email})
    candidate_passwords = [payload.password]
    trimmed_password = payload.password.strip()
    if trimmed_password and trimmed_password != payload.password:
        candidate_passwords.append(trimmed_password)
    password_ok = bool(user) and any(
        verify_password(candidate, user["password_hash"])
        for candidate in candidate_passwords
    )
    if not password_ok:
        now = datetime.now(timezone.utc)
        await db.login_attempts.insert_one(
            {
                "identifier": identifier,
                "email": email,
                "at": now.isoformat(),
                "expires_at": now + timedelta(minutes=LOCKOUT_MINUTES * 2),
            })
        raise HTTPException(status_code=401, detail="E-mail ou senha inválidos")

    await db.login_attempts.delete_many({"identifier": identifier})
    uid = str(user["_id"])
    ver = user.get("token_version", 0)
    now_iso = datetime.now(timezone.utc).isoformat()
    await db.users.update_one({"_id": user["_id"]}, {"$set": {"last_login_at": now_iso}})
    if user.get("role", "admin") == "admin":
        await db.system_settings.update_one(
            {"_id": "ticket_notification_admin"},
            {"$set": {
                "user_id": uid,
                "email": str(user.get("email") or "").strip().lower(),
                "name": user.get("name", ""),
                "updated_at": now_iso,
            }},
            upsert=True,
        )
    set_auth_cookies(response, create_access_token(uid, email, ver), create_refresh_token(uid, ver))
    return {"id": uid, "email": user["email"], "name": user.get("name", "Admin"), "role": user.get("role", "admin")}


@api_router.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    response.delete_cookie("refresh_token", path="/")
    return {"message": "Sessão encerrada"}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return {"id": user["_id"], "email": user["email"], "name": user.get("name", "Admin"), "role": user.get("role", "admin")}


@api_router.get("/users")
async def list_users(user: dict = Depends(require_admin)):
    docs = await db.users.find().sort("created_at", 1).to_list(500)
    return [ser_user(d) for d in docs]


@api_router.post("/users")
async def create_user(payload: UserCreate, background_tasks: BackgroundTasks, user: dict = Depends(require_admin)):
    email = payload.email.lower().strip()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="E-mail já cadastrado")
    if payload.password:
        if len(payload.password) < MIN_PASSWORD_LENGTH:
            raise HTTPException(status_code=400, detail=f"A senha deve ter ao menos {MIN_PASSWORD_LENGTH} caracteres")
        pw_hash = hash_password(payload.password)
    else:
        pw_hash = hash_password(secrets.token_urlsafe(16))
    role = payload.role if payload.role in ("admin", "responsavel") else "responsavel"
    categories = [] if role == "admin" else await validate_category_ids(payload.categories)
    doc = {
        "email": email, "password_hash": pw_hash,
        "name": payload.name.strip() or "Responsável", "role": role,
        "categories": categories,
        "token_version": 0, "created_at": datetime.now(timezone.utc).isoformat(),
    }
    res = await db.users.insert_one(doc)
    doc["_id"] = res.inserted_id
    await write_audit(user, "user.create", "user", str(res.inserted_id), {
        "email": email,
        "role": role,
        "categories_count": len(doc.get("categories") or []),
    })
    if payload.send_welcome:
        token = secrets.token_urlsafe(32)
        await db.password_reset_tokens.insert_one({
            "token_hash": hashlib.sha256(token.encode()).hexdigest(),
            "user_id": str(res.inserted_id), "email": email,
            "expires_at": (datetime.now(timezone.utc) + timedelta(hours=48)).isoformat(),
            "used": False, "created_at": datetime.now(timezone.utc).isoformat(),
            "purge_at": datetime.now(timezone.utc) + timedelta(hours=72),
        })
        background_tasks.add_task(send_welcome_email, email, doc["name"], token)
    return ser_user(doc)


@api_router.put("/users/{uid}")
async def update_user(uid: str, payload: UserUpdate, user: dict = Depends(require_admin)):
    oid = as_object_id(uid, "Usuário não encontrado")
    target = await db.users.find_one({"_id": oid})
    if not target:
        raise HTTPException(status_code=404, detail="Usuário não encontrado")
    updates = {}
    if payload.name is not None:
        cleaned_name = payload.name.strip()
        if not cleaned_name:
            raise HTTPException(status_code=400, detail="O nome do usuário é obrigatório")
        updates["name"] = cleaned_name

    next_role = payload.role or target.get("role", "responsavel")
    if target.get("role") == "admin" and next_role != "admin":
        admin_count = await db.users.count_documents({"role": "admin"})
        if admin_count <= 1:
            raise HTTPException(status_code=400, detail="Deve existir ao menos um administrador")

    if payload.role is not None:
        updates["role"] = next_role

    if next_role == "admin":
        if payload.categories is not None or payload.role == "admin":
            updates["categories"] = []
    elif payload.categories is not None:
        updates["categories"] = await validate_category_ids(payload.categories)

    inc = {}
    if payload.password:
        if len(payload.password) < MIN_PASSWORD_LENGTH:
            raise HTTPException(status_code=400, detail=f"A senha deve ter ao menos {MIN_PASSWORD_LENGTH} caracteres")
        updates["password_hash"] = hash_password(payload.password)
        inc["token_version"] = 1
    if updates or inc:
        op = {}
        if updates:
            op["$set"] = updates
        if inc:
            op["$inc"] = inc
        await db.users.update_one({"_id": oid}, op)
    updated = await db.users.find_one({"_id": oid})
    await write_audit(user, "user.update", "user", uid, {
        "email": updated.get("email", ""),
        "role": updated.get("role", ""),
        "categories_count": len(updated.get("categories") or []),
        "password_changed": bool(payload.password),
    })
    return ser_user(updated)


@api_router.delete("/users/{uid}")
async def delete_user(uid: str, user: dict = Depends(require_admin)):
    oid = as_object_id(uid, "Usuário não encontrado")
    if str(user["_id"]) == uid:
        raise HTTPException(status_code=400, detail="Você não pode remover seu próprio usuário")
    if await db.users.count_documents({}) <= 1:
        raise HTTPException(status_code=400, detail="Deve existir ao menos um usuário")
    if target.get("role") == "admin":
        admin_count = await db.users.count_documents({"role": "admin"})
        if admin_count <= 1:
            raise HTTPException(status_code=400, detail="Deve existir ao menos um administrador")
    await db.users.delete_one({"_id": oid})
    await write_audit(user, "user.delete", "user", uid, {
        "email": target.get("email", ""),
        "name": target.get("name", ""),
    })
    return {"message": "Usuário removido"}


@api_router.post("/auth/refresh")
async def refresh(request: Request, response: Response):
    token = request.cookies.get("refresh_token")
    if not token:
        raise HTTPException(status_code=401, detail="Sem token de atualização")
    try:
        payload = jwt.decode(token, get_jwt_secret(), algorithms=[JWT_ALGORITHM])
        if payload.get("type") != "refresh":
            raise HTTPException(status_code=401, detail="Tipo de token inválido")
        user = await db.users.find_one({"_id": as_object_id(payload["sub"], "Sessão inválida")})
        if not user or payload.get("ver", 0) != user.get("token_version", 0):
            raise HTTPException(status_code=401, detail="Sessão expirada")
        uid = str(user["_id"])
        ver = user.get("token_version", 0)
        set_auth_cookies(response, create_access_token(uid, user["email"], ver),
                         create_refresh_token(uid, ver))
        return {"message": "ok"}
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token inválido")


GENERIC_RESET_MSG = {"message": "Se este e-mail estiver cadastrado, enviaremos um link de redefinição."}


@api_router.post("/auth/forgot-password")
async def forgot_password(payload: ForgotInput, request: Request, background_tasks: BackgroundTasks):
    await enforce_rate_limit(request, "forgot-password", limit=5, minutes=15)
    email = payload.email.lower().strip()
    now = datetime.now(timezone.utc)
    await db.password_reset_requests.insert_one({
        "email": email,
        "created_at": now.isoformat(),
        "expires_at": now + timedelta(hours=24),
    })
    window = (now - timedelta(minutes=15)).isoformat()
    recent = await db.password_reset_requests.count_documents(
        {"email": email, "created_at": {"$gt": window}})
    if recent > 5:
        return GENERIC_RESET_MSG
    user = await db.users.find_one({"email": email})
    if not user:
        return GENERIC_RESET_MSG
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    await db.password_reset_tokens.insert_one({
        "token_hash": token_hash, "user_id": str(user["_id"]), "email": email,
        "expires_at": (now + timedelta(hours=1)).isoformat(), "used": False,
        "created_at": now.isoformat(),
        "purge_at": now + timedelta(hours=24),
    })
    background_tasks.add_task(send_password_reset_email, user["email"], token)
    return GENERIC_RESET_MSG


@api_router.post("/auth/reset-password")
async def reset_password(payload: ResetInput):
    if len(payload.password) < MIN_PASSWORD_LENGTH:
        raise HTTPException(status_code=400, detail=f"A senha deve ter ao menos {MIN_PASSWORD_LENGTH} caracteres")
    now_iso = datetime.now(timezone.utc).isoformat()
    h = hashlib.sha256(payload.token.encode()).hexdigest()
    doc = await db.password_reset_tokens.find_one_and_update(
        {"token_hash": h, "used": False, "expires_at": {"$gt": now_iso}},
        {"$set": {"used": True}})
    if not doc:
        raise HTTPException(status_code=400, detail="Link inválido, expirado ou já utilizado")
    await db.users.update_one(
        {"_id": as_object_id(doc["user_id"], "Usuário não encontrado")},
        {"$set": {"password_hash": hash_password(payload.password)}, "$inc": {"token_version": 1}})
    await db.password_reset_tokens.delete_many({"user_id": doc["user_id"], "used": False})
    await db.login_attempts.delete_many({"email": doc["email"]})
    return {"message": "Senha redefinida com sucesso"}


# ----------------------------------------------------------------------------
# Category endpoints
# ----------------------------------------------------------------------------
register_material_routes(api_router, lambda: db, require_admin, write_audit)


@api_router.get("/categories")
async def list_categories(request: Request, all: bool = False):
    if all:
        user = await get_current_user(request)
        if user.get("role") == "admin":
            docs = await db.categories.find({}).sort("created_at", 1).to_list(200)
            return [ser_category(d) for d in docs]

        assigned_ids = [
            ObjectId(value)
            for value in (user.get("categories") or [])
            if ObjectId.is_valid(str(value))
        ]
        if not assigned_ids:
            return []
        docs = await db.categories.find({
            "_id": {"$in": assigned_ids},
            "active": True,
        }).sort("created_at", 1).to_list(200)
        return [ser_category_public(d) for d in docs]

    docs = await db.categories.find({"active": True}).sort("created_at", 1).to_list(200)
    return [ser_category_public(d) for d in docs]


@api_router.get("/categories/{cat_id}")
async def get_category(cat_id: str):
    try:
        oid = ObjectId(cat_id)
    except Exception:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    doc = await db.categories.find_one({"_id": oid, "active": True})
    if not doc:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    return ser_category_public(doc)


@api_router.get("/categories/{cat_id}/template")
async def category_template(cat_id: str):
    try:
        oid = ObjectId(cat_id)
    except Exception:
        raise HTTPException(status_code=404, detail="Modelo não disponível")
    doc = await db.categories.find_one({"_id": oid, "active": True})
    cols = (doc or {}).get("template_columns") or []
    if not doc or not cols:
        raise HTTPException(status_code=404, detail="Modelo não disponível")
    wb = Workbook()
    ws = wb.active
    ws.title = "Modelo"
    ws.append(cols)
    for i, _ in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="660099")
        ws.column_dimensions[cell.column_letter].width = 30
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = os.path.basename(doc.get("template_filename") or "modelo.xlsx").replace("\r", "").replace("\n", "").replace('"', "")
    return StarletteResponse(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"',
                 "X-Content-Type-Options": "nosniff"},
    )


@api_router.post("/categories")
async def create_category(payload: CategoryInput, user: dict = Depends(require_admin)):
    normalized_name = payload.name.strip()
    duplicate = await db.categories.find_one({
        "name": {"$regex": f"^{re.escape(normalized_name)}$", "$options": "i"}
    })
    if duplicate:
        raise HTTPException(status_code=409, detail="Já existe uma categoria com este nome")
    kit_catalog_ids = await validate_kit_catalog_ids(payload.kit_catalog_ids)
    if payload.kit_enabled and not kit_catalog_ids:
        raise HTTPException(status_code=400, detail="Vincule ao menos um catálogo para habilitar Solicitação de Kit")
    doc = payload.model_dump()
    doc["name"] = normalized_name
    doc["owners"] = [str(owner).lower() for owner in payload.owners]
    doc["kit_catalog_ids"] = kit_catalog_ids
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    res = await db.categories.insert_one(doc)
    doc["_id"] = res.inserted_id
    await write_audit(user, "category.create", "category", str(res.inserted_id), {
        "name": normalized_name,
        "active": doc.get("active", True),
    })
    return ser_category(doc)


@api_router.put("/categories/{cat_id}")
async def update_category(cat_id: str, payload: CategoryInput, user: dict = Depends(require_admin)):
    oid = as_object_id(cat_id, "Categoria não encontrada")
    normalized_name = payload.name.strip()
    duplicate = await db.categories.find_one({
        "_id": {"$ne": oid},
        "name": {"$regex": f"^{re.escape(normalized_name)}$", "$options": "i"},
    })
    if duplicate:
        raise HTTPException(status_code=409, detail="Já existe uma categoria com este nome")
    kit_catalog_ids = await validate_kit_catalog_ids(payload.kit_catalog_ids)
    if payload.kit_enabled and not kit_catalog_ids:
        raise HTTPException(status_code=400, detail="Vincule ao menos um catálogo para habilitar Solicitação de Kit")
    doc = payload.model_dump()
    doc["name"] = normalized_name
    doc["owners"] = [str(owner).lower() for owner in payload.owners]
    doc["kit_catalog_ids"] = kit_catalog_ids
    res = await db.categories.update_one({"_id": oid}, {"$set": doc})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    updated = await db.categories.find_one({"_id": oid})
    await write_audit(user, "category.update", "category", cat_id, {
        "name": updated.get("name", ""),
        "active": updated.get("active", True),
    })
    return ser_category(updated)


@api_router.delete("/categories/{cat_id}")
async def delete_category(cat_id: str, user: dict = Depends(require_admin)):
    oid = as_object_id(cat_id, "Categoria não encontrada")
    target = await db.categories.find_one({"_id": oid})
    if not target:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    result = await db.categories.delete_one({"_id": oid})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Categoria não encontrada")
    await db.users.update_many({}, {"$pull": {"categories": cat_id}})
    await db.categories.update_many({"kit_catalog_ids": cat_id}, {"$pull": {"kit_catalog_ids": cat_id}})
    await db.categories.update_many(
        {"kit_enabled": True, "kit_catalog_ids": []},
        {"$set": {"kit_enabled": False}},
    )
    await write_audit(user, "category.delete", "category", cat_id, {
        "name": target.get("name", ""),
    })
    return {"message": "Categoria removida"}


# ----------------------------------------------------------------------------
# Ticket endpoints
# ----------------------------------------------------------------------------
async def next_ticket_number() -> str:
    year = datetime.now(timezone.utc).year
    res = await db.counters.find_one_and_update(
        {"_id": f"tickets-{year}"}, {"$inc": {"seq": 1}},
        upsert=True, return_document=True)
    seq = res["seq"] if res else 1
    return f"CH-{year}-{seq:04d}"


@api_router.post("/tickets")
async def create_ticket(
    request: Request,
    background_tasks: BackgroundTasks,
    payload: str = Form(...),
    file: Optional[UploadFile] = File(None),
):
    await enforce_rate_limit(request, "ticket-create", limit=10, minutes=1)
    if len(payload.encode("utf-8")) > 256 * 1024:
        raise HTTPException(status_code=413, detail="Solicitação muito grande")
    try:
        data = json.loads(payload)
    except Exception:
        raise HTTPException(status_code=400, detail="Dados da solicitação inválidos")

    if not isinstance(data, dict):
        raise HTTPException(400, "Dados da solicitação inválidos")
    kit_mode = data.get("kind") == "kit"
    material_items, kit_categories = [], []
    if kit_mode or data.get("material_items"):
        material_items, kit_categories = await validate_kit(db, data.get("material_items"))
    if kit_mode:
        category = {
            "_id": kit_categories[0]["_id"], "name": "Kit de materiais", "icon": "Package",
            "lead_time_hours": max(int(c.get("lead_time_hours", 24)) for c in kit_categories),
            "fields": [{**f, "label": f"{c['name']} — {f['label']}"} for c in kit_categories for f in c.get("fields", [])],
        }
    else:
        category_id = str(data.get("category_id") or "").strip()
        if not category_id:
            raise HTTPException(status_code=400, detail="Categoria obrigatória")
        category = await db.categories.find_one({"_id": as_object_id(category_id, "Categoria não encontrada"), "active": True})
        if not category:
            raise HTTPException(status_code=404, detail="Categoria não encontrada")

        if category.get("kit_enabled"):
            allowed_catalog_ids = set(category.get("kit_catalog_ids") or [])
            if not allowed_catalog_ids:
                raise HTTPException(400, "Esta categoria está sem catálogo de kit vinculado")
            if not material_items:
                raise HTTPException(400, "Selecione pelo menos um sub-item para montar o kit")
            if any(item["category_id"] not in allowed_catalog_ids for item in material_items):
                raise HTTPException(400, "Um dos materiais não pertence aos catálogos vinculados a esta categoria")
        else:
            if any(item["category_id"] != category_id for item in material_items):
                raise HTTPException(400, "Para combinar categorias diferentes, use a montagem de kit")
            if any(item.get("active", True) for item in category.get("materials", [])) and not material_items:
                raise HTTPException(400, "Selecione pelo menos um sub-item")

    requester = data.get("requester", {})
    if not isinstance(requester, dict):
        raise HTTPException(400, "Dados do solicitante inválidos")
    for field in ("matricula", "email", "empresa"):
        if not isinstance(requester.get(field), str) or not requester[field].strip():
            raise HTTPException(status_code=400, detail=f"Campo obrigatório: {field}")
    requester_email = str(requester.get("email", "")).strip().lower()
    if len(requester_email) > 254 or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", requester_email):
        raise HTTPException(status_code=400, detail="Informe um e-mail válido")
    if len(str(requester.get("matricula", ""))) > 80 or len(str(requester.get("empresa", ""))) > 160:
        raise HTTPException(status_code=400, detail="Dados do solicitante excedem o limite permitido")

    field_values = data.get("field_values") or {}
    if not isinstance(field_values, dict) or len(field_values) > (1000 if kit_mode else 100):
        raise HTTPException(status_code=400, detail="Campos da solicitação inválidos")
    for field in category.get("fields", []):
        label = field.get("label")
        field_type = field.get("type")
        value = field_values.get(label)

        if field_type == "dependent_select":
            pair = value if isinstance(value, dict) else {}
            selected_parent = str(pair.get("parent") or "").strip()
            raw_children = pair.get("children")
            if isinstance(raw_children, list):
                selected_children = [
                    str(child).strip()
                    for child in raw_children
                    if str(child).strip()
                ]
            else:
                legacy_child = str(pair.get("child") or "").strip()
                selected_children = [legacy_child] if legacy_child else []

            if field.get("required") and (not selected_parent or not selected_children):
                raise HTTPException(status_code=400, detail=f"Campo obrigatório: {label or 'campo'}")

            if selected_parent or selected_children:
                allowed_pairs = {
                    str(item.get("parent") or "").strip(): [
                        str(child).strip()
                        for child in (item.get("children") or [])
                        if str(child).strip()
                    ]
                    for item in (field.get("dependent_options") or [])
                    if str(item.get("parent") or "").strip()
                }
                allowed_children = allowed_pairs.get(selected_parent, [])
                if (
                    not selected_parent
                    or not selected_children
                    or selected_parent not in allowed_pairs
                    or any(child not in allowed_children for child in selected_children)
                ):
                    raise HTTPException(
                        status_code=400,
                        detail=f"Seleção inválida no campo: {label or 'campo'}",
                    )

                # Normaliza o formato novo e remove seleções duplicadas mantendo a ordem.
                field_values[label] = {
                    "parent": selected_parent,
                    "children": list(dict.fromkeys(selected_children)),
                }
            continue

        if not field.get("required"):
            continue
        missing = value is not True if field_type == "checkbox" else not str(value or "").strip()
        if missing:
            raise HTTPException(status_code=400, detail=f"Campo obrigatório: {label or 'campo'}")
    for key, value in field_values.items():
        if len(str(key)) > (285 if kit_mode else 160) or len(str(value)) > 10000:
            raise HTTPException(status_code=400, detail="Um dos campos excede o limite permitido")

    file_ref = None
    if file is not None:
        raw = await file.read()
        if raw:
            if len(raw) > 10 * 1024 * 1024:
                raise HTTPException(status_code=400, detail="Arquivo excede 10MB")
            safe_name = os.path.basename(file.filename or "arquivo").replace("\r", "").replace("\n", "")
            ext = safe_name.rsplit(".", 1)[-1].lower() if "." in safe_name else ""
            allowed_extensions = {"pdf", "png", "jpg", "jpeg", "doc", "docx", "xls", "xlsx", "csv", "txt"}
            if ext not in allowed_extensions:
                raise HTTPException(status_code=400, detail="Tipo de arquivo não permitido")
            path = f"{APP_NAME}/tickets/{uuid.uuid4()}.{ext}"
            result = await put_object(path, raw, file.content_type or "application/octet-stream")
            file_ref = {
                "storage_path": result["path"],
                "original_filename": safe_name,
                "content_type": file.content_type or "application/octet-stream",
                "size": result.get("size", len(raw)),
            }

    lead = int(category.get("lead_time_hours", 24))
    now = datetime.now(timezone.utc)
    number = await next_ticket_number()
    ticket = {
        "ticket_number": number,
        "category_id": str(category["_id"]),
        "category_name": category.get("name", ""),
        "category_icon": category.get("icon", "Laptop"),
        "requester": {
            "matricula": requester["matricula"].strip(),
            "email": requester["email"].strip().lower(),
            "empresa": requester["empresa"].strip(),
        },
        "field_values": field_values,
        "material_items": material_items,
        "category_ids": [str(c["_id"]) for c in kit_categories] if kit_mode else [str(category["_id"])],
        "kind": "kit" if kit_mode else "standard",
        "file": file_ref,
        "status": "aberto",
        "lead_time_hours": lead,
        "due_at": (now + timedelta(hours=lead)).isoformat(),
        "created_at": now.isoformat(),
        "history": [{"status": "aberto", "at": now.isoformat(), "note": "Chamado aberto"}],
        "email_log": [],
    }
    res = await db.tickets.insert_one(ticket)
    ticket["_id"] = res.inserted_id

    owners = await get_all_admin_notification_emails()
    if owners:
        background_tasks.add_task(notify_owners, dict(ticket), owners)
    background_tasks.add_task(send_ticket_opened_email, dict(ticket))

    return ser_ticket(ticket)


@api_router.post("/email/brevo-webhook")
async def brevo_email_webhook(request: Request):
    if not BREVO_WEBHOOK_SECRET:
        raise HTTPException(status_code=503, detail="Webhook não configurado")
    supplied = request.headers.get("x-brevo-webhook-secret") or ""
    if not secrets.compare_digest(supplied, BREVO_WEBHOOK_SECRET):
        raise HTTPException(status_code=401, detail="Webhook não autorizado")
    payload = await request.json()
    events = payload if isinstance(payload, list) else [payload]
    now = datetime.now(timezone.utc).isoformat()
    stored = 0
    for event in events:
        if not isinstance(event, dict):
            continue
        email = str(event.get("email") or "").strip().lower()
        event_name = str(event.get("event") or "").strip()
        message_id = str(event.get("message-id") or event.get("messageId") or "").strip()
        doc = {
            "provider": "brevo",
            "event": event_name,
            "email": email,
            "message_id": message_id,
            "subject": str(event.get("subject") or ""),
            "reason": str(event.get("reason") or ""),
            "raw_date": str(event.get("date") or ""),
            "received_at": now,
        }
        await db.email_events.insert_one(doc)
        stored += 1
        logger.info(
            "Brevo delivery event=%s email=%s message_id=%s",
            event_name or "unknown",
            email or "unknown",
            message_id or "unknown",
        )
    return {"ok": True, "stored": stored}


@api_router.get("/tickets/track")
async def track_ticket(q: str, request: Request):
    await enforce_rate_limit(request, "ticket-track", limit=30, minutes=1)
    q = q.strip()
    if not q or len(q) > 254:
        return []
    query = {"$or": [
        {"ticket_number": q.upper()},
        {"requester.email": q.lower()},
        {"requester.matricula": q},
    ]}
    docs = await db.tickets.find(query).sort("created_at", -1).to_list(50)
    return [ser_ticket_public(d) for d in docs]


def build_ticket_query(
    user: dict,
    status: Optional[str] = None,
    category_id: Optional[str] = None,
    search: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
) -> dict:
    query = {}
    if status and status != "all":
        if status not in STATUS_LABELS:
            raise HTTPException(status_code=400, detail="Status inválido")
        query["status"] = status
    if category_id and category_id != "all":
        query.setdefault("$and", []).append(ticket_category_match([category_id]))
    if search:
        escaped = re.escape(search.strip()[:160])
        query["$or"] = [
            {"ticket_number": {"$regex": escaped, "$options": "i"}},
            {"requester.email": {"$regex": escaped, "$options": "i"}},
            {"requester.matricula": {"$regex": escaped, "$options": "i"}},
            {"requester.empresa": {"$regex": escaped, "$options": "i"}},
        ]

    if start_date or end_date:
        created_range = {}
        try:
            if start_date:
                start = datetime.strptime(start_date, "%Y-%m-%d").replace(tzinfo=timezone.utc)
                created_range["$gte"] = start.isoformat()
            if end_date:
                end = datetime.strptime(end_date, "%Y-%m-%d").replace(tzinfo=timezone.utc) + timedelta(days=1)
                created_range["$lt"] = end.isoformat()
        except ValueError:
            raise HTTPException(status_code=400, detail="Período do relatório inválido")
        query["created_at"] = created_range

    if user.get("role") != "admin":
        cats = user.get("categories", []) or []
        if category_id and category_id != "all" and category_id not in cats:
            query.setdefault("$and", []).append(ticket_category_match([]))
        query.setdefault("$and", []).append(ticket_category_match(cats))
    return query


@api_router.get("/tickets")
async def list_tickets(status: Optional[str] = None, category_id: Optional[str] = None,
                       search: Optional[str] = None, start_date: Optional[str] = None,
                       end_date: Optional[str] = None, page: int = 1, limit: int = 50,
                       paginated: bool = False, user: dict = Depends(get_current_user)):
    query = build_ticket_query(user, status, category_id, search, start_date, end_date)

    if paginated:
        safe_limit = min(max(limit, 10), 100)
        safe_page = max(page, 1)
        total = await db.tickets.count_documents(query)
        docs = await (
            db.tickets.find(query)
            .sort("created_at", -1)
            .skip((safe_page - 1) * safe_limit)
            .limit(safe_limit)
            .to_list(safe_limit)
        )
        pages = max(1, (total + safe_limit - 1) // safe_limit)
        return {
            "items": [ser_ticket(d) for d in docs],
            "total": total,
            "page": safe_page,
            "limit": safe_limit,
            "pages": pages,
        }

    docs = await db.tickets.find(query).sort("created_at", -1).to_list(1000)
    return [ser_ticket(d) for d in docs]


@api_router.get("/admin/reports/tickets.csv")
async def export_ticket_report(
    status: Optional[str] = None,
    category_id: Optional[str] = None,
    search: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    user: dict = Depends(get_current_user),
):
    query = build_ticket_query(user, status, category_id, search, start_date, end_date)
    docs = await db.tickets.find(query).sort("created_at", -1).limit(10000).to_list(10000)

    output = StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "Chamado", "Categoria", "Status", "Matrícula", "E-mail",
        "Empresa", "Aberto em", "Prazo (h)", "Previsão", "Materiais do kit"
    ])
    for doc in docs:
        requester = doc.get("requester") or {}
        writer.writerow([
            doc.get("ticket_number", ""),
            doc.get("category_name", ""),
            STATUS_LABELS.get(doc.get("status"), doc.get("status", "")),
            requester.get("matricula", ""),
            requester.get("email", ""),
            requester.get("empresa", ""),
            doc.get("created_at", ""),
            doc.get("lead_time_hours", ""),
            doc.get("due_at", ""),
            " | ".join(f"{i['category_name']} / {i['name']}: {i['quantity']} {'m' if i['measure'] == 'metro' else 'un'}" for i in doc.get("material_items", [])),
        ])

    await write_audit(user, "report.export", "report", "tickets", {
        "rows": len(docs),
        "status": status or "all",
        "category_id": category_id or "all",
        "start_date": start_date or "",
        "end_date": end_date or "",
    })
    filename = f"relatorio-chamados-{datetime.now(timezone.utc).date().isoformat()}.csv"
    return StarletteResponse(
        content=output.getvalue().encode("utf-8-sig"),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "X-Content-Type-Options": "nosniff",
        },
    )


@api_router.get("/tickets/{ticket_id}")
async def get_ticket(ticket_id: str, user: dict = Depends(get_current_user)):
    oid = as_object_id(ticket_id, "Chamado não encontrado")
    doc = await db.tickets.find_one({"_id": oid})
    if not doc:
        raise HTTPException(status_code=404, detail="Chamado não encontrado")
    if not can_access_ticket(user, doc):
        raise HTTPException(status_code=404, detail="Chamado não encontrado")
    return ser_ticket(doc)


@api_router.patch("/tickets/{ticket_id}/status")
async def update_status(
    ticket_id: str,
    payload: StatusUpdate,
    background_tasks: BackgroundTasks,
    user: dict = Depends(get_current_user),
):
    if payload.status not in STATUS_LABELS:
        raise HTTPException(status_code=400, detail="Status inválido")
    oid = as_object_id(ticket_id, "Chamado não encontrado")
    doc = await db.tickets.find_one({"_id": oid})
    if not doc:
        raise HTTPException(status_code=404, detail="Chamado não encontrado")
    if not can_access_ticket(user, doc):
        raise HTTPException(status_code=404, detail="Chamado não encontrado")
    now = datetime.now(timezone.utc).isoformat()
    actor = user.get("name") or user.get("email")
    responder = {
        "id": str(user.get("_id") or ""),
        "name": str(user.get("name") or user.get("email") or "Equipe de atendimento"),
        "email": str(user.get("email") or "").strip().lower(),
    }
    entry = {
        "status": payload.status,
        "at": now,
        "note": payload.note or "",
        "by": actor,
        "by_email": responder["email"],
    }
    previous_status = doc.get("status")
    set_fields = {"status": payload.status}
    if payload.status in {"em_analise", "em_andamento"}:
        set_fields["assigned_to"] = responder
        set_fields["assigned_at"] = now
    await db.tickets.update_one(
        {"_id": oid},
        {"$set": set_fields, "$push": {"history": entry}})
    doc = await db.tickets.find_one({"_id": oid})
    await write_audit(user, "ticket.status", "ticket", ticket_id, {
        "ticket_number": doc.get("ticket_number", ""),
        "from": previous_status,
        "to": payload.status,
    })
    background_tasks.add_task(
        send_status_update_email,
        dict(doc),
        previous_status,
        payload.status,
        payload.note,
        responder if payload.status in {"em_analise", "em_andamento"} else (doc.get("assigned_to") or responder),
    )
    return ser_ticket(doc)


@api_router.get("/files/{path:path}")
async def download_file(path: str, user: dict = Depends(get_current_user)):
    record = await db.tickets.find_one({"file.storage_path": path})
    if not record:
        raise HTTPException(status_code=404, detail="Arquivo não encontrado")
    if not can_access_ticket(user, record):
        raise HTTPException(status_code=404, detail="Arquivo não encontrado")
    data, content_type = await get_object(path)
    fname = os.path.basename(record["file"].get("original_filename", "arquivo")).replace("\r", "").replace("\n", "").replace('"', "")
    return StarletteResponse(content=data, media_type=record["file"].get("content_type", content_type),
                             headers={"Content-Disposition": f'attachment; filename="{fname}"',
                                      "X-Content-Type-Options": "nosniff"})


# ----------------------------------------------------------------------------
# Admin audit
# ----------------------------------------------------------------------------
@api_router.get("/admin/audit")
async def admin_audit(page: int = 1, limit: int = 50, user: dict = Depends(require_admin)):
    safe_limit = min(max(limit, 10), 100)
    safe_page = max(page, 1)
    total = await db.audit_logs.count_documents({})
    docs = await (
        db.audit_logs.find({})
        .sort("created_at", -1)
        .skip((safe_page - 1) * safe_limit)
        .limit(safe_limit)
        .to_list(safe_limit)
    )
    items = []
    for doc in docs:
        item = dict(doc)
        item["id"] = str(item.pop("_id"))
        items.append(item)
    return {
        "items": items,
        "total": total,
        "page": safe_page,
        "pages": max(1, (total + safe_limit - 1) // safe_limit),
        "limit": safe_limit,
    }


# ----------------------------------------------------------------------------
# Admin stats
# ----------------------------------------------------------------------------
@api_router.get("/admin/stats")
async def admin_stats(user: dict = Depends(get_current_user)):
    match = scope_match(user)
    active_statuses = ["aberto", "em_analise", "em_andamento"]
    now = datetime.now(timezone.utc)
    now_iso = now.isoformat()
    due_soon_iso = (now + timedelta(hours=6)).isoformat()
    total = await db.tickets.count_documents(match)
    open_count = await db.tickets.count_documents({**match, "status": {"$in": active_statuses}})
    done = await db.tickets.count_documents({**match, "status": "concluido"})
    overdue = await db.tickets.count_documents({
        **match,
        "status": {"$in": active_statuses},
        "due_at": {"$lt": now_iso},
    })
    due_soon = await db.tickets.count_documents({
        **match,
        "status": {"$in": active_statuses},
        "due_at": {"$gte": now_iso, "$lte": due_soon_iso},
    })
    by_cat = await db.tickets.aggregate([
        {"$match": match},
        {"$group": {"_id": "$category_name", "count": {"$sum": 1}}},
        {"$sort": {"count": -1}}
    ]).to_list(20)
    by_status = await db.tickets.aggregate([
        {"$match": match},
        {"$group": {"_id": "$status", "count": {"$sum": 1}}}
    ]).to_list(20)
    recent_docs = await (
        db.tickets.find(
            match,
            {
                "ticket_number": 1,
                "category_name": 1,
                "category_icon": 1,
                "status": 1,
                "created_at": 1,
            },
        )
        .sort("created_at", -1)
        .limit(6)
        .to_list(6)
    )
    recent = [
        {
            "id": str(item["_id"]),
            "ticket_number": item.get("ticket_number", ""),
            "category_name": item.get("category_name", ""),
            "category_icon": item.get("category_icon", "CircleHelp"),
            "status": item.get("status", "aberto"),
            "status_label": STATUS_LABELS.get(item.get("status"), item.get("status")),
            "created_at": item.get("created_at"),
        }
        for item in recent_docs
    ]
    return {
        "total": total,
        "open": open_count,
        "done": done,
        "overdue": overdue,
        "due_soon": due_soon,
        "recent": recent,
        "by_category": [{"name": c["_id"], "count": c["count"]} for c in by_cat],
        "by_status": [{"status": s["_id"], "label": STATUS_LABELS.get(s["_id"], s["_id"]), "count": s["count"]} for s in by_status],
    }


# ----------------------------------------------------------------------------
# Seeding
# ----------------------------------------------------------------------------
DEFAULT_CATEGORIES = [
    {
        "name": "Informática & TI", "icon": "Laptop",
        "description": "Suporte a computadores, periféricos, softwares, rede e sistemas de TI",
        "lead_time_hours": 4, "owners": DEFAULT_OWNERS, "active": True,
        "fields": [
            {"id": str(uuid.uuid4()), "label": "Tipo de problema", "type": "select", "required": True,
             "options": ["Hardware", "Software", "Rede/Internet", "Impressora", "E-mail", "Outro"]},
            {"id": str(uuid.uuid4()), "label": "Equipamento / Patrimônio", "type": "text", "required": False, "options": []},
            {"id": str(uuid.uuid4()), "label": "Descreva o problema", "type": "textarea", "required": True, "options": []},
            {"id": str(uuid.uuid4()), "label": "É urgente?", "type": "checkbox", "required": False, "options": []},
        ],
    },
    {
        "name": "Acessos e Permissões", "icon": "KeyRound",
        "description": "Criação, alteração ou revogação de acessos a sistemas e pastas compartilhadas",
        "lead_time_hours": 8, "owners": DEFAULT_OWNERS, "active": True,
        "fields": [
            {"id": str(uuid.uuid4()), "label": "Sistema / Recurso", "type": "text", "required": True, "options": []},
            {"id": str(uuid.uuid4()), "label": "Tipo de solicitação", "type": "select", "required": True,
             "options": ["Criar acesso", "Alterar acesso", "Revogar acesso"]},
            {"id": str(uuid.uuid4()), "label": "Justificativa", "type": "textarea", "required": True, "options": []},
        ],
    },
    {
        "name": "Férias e RH", "icon": "CalendarCheck",
        "description": "Solicitação de agendamento, alteração ou dúvidas sobre férias e benefícios",
        "lead_time_hours": 24, "owners": DEFAULT_OWNERS, "active": True,
        "fields": [
            {"id": str(uuid.uuid4()), "label": "Tipo de solicitação", "type": "select", "required": True,
             "options": ["Agendar férias", "Alterar férias", "Dúvida sobre benefícios"]},
            {"id": str(uuid.uuid4()), "label": "Data de início desejada", "type": "date", "required": False, "options": []},
            {"id": str(uuid.uuid4()), "label": "Observações", "type": "textarea", "required": False, "options": []},
        ],
    },
    {
        "name": "Criação de Centros", "icon": "Building2",
        "description": "Abertura de novos centros de custo, unidades operacionais ou projetos",
        "lead_time_hours": 48, "owners": DEFAULT_OWNERS, "active": True,
        "template_columns": ["Endereço físico", "CNPJ", "Inscrição Estadual"],
        "template_filename": "modelo-criacao-centros.xlsx",
        "fields": [
            {"id": str(uuid.uuid4()), "label": "Nome do centro", "type": "text", "required": True, "options": []},
            {"id": str(uuid.uuid4()), "label": "Tipo de centro", "type": "select", "required": True,
             "options": ["Centro de custo", "Unidade operacional", "Projeto"]},
            {"id": str(uuid.uuid4()), "label": "Responsável do centro", "type": "text", "required": True, "options": []},
            {"id": str(uuid.uuid4()), "label": "Detalhes", "type": "textarea", "required": False, "options": []},
        ],
    },
    {
        "name": "Reabastecimento", "icon": "PackagePlus",
        "description": "Solicitação de materiais de escritório, insumos de copa e suprimentos",
        "lead_time_hours": 12, "owners": DEFAULT_OWNERS, "active": True,
        "fields": [
            {"id": str(uuid.uuid4()), "label": "Categoria do item", "type": "select", "required": True,
             "options": ["Material de escritório", "Copa/Cozinha", "Limpeza", "Suprimentos de TI"]},
            {"id": str(uuid.uuid4()), "label": "Itens e quantidades", "type": "textarea", "required": True, "options": []},
            {"id": str(uuid.uuid4()), "label": "Local de entrega", "type": "text", "required": True, "options": []},
        ],
    },
]


async def seed_admin():
    admin_email = (os.environ.get("ADMIN_EMAIL") or "").strip().lower()
    admin_password = os.environ.get("ADMIN_PASSWORD") or ""
    admin_password_hash = (os.environ.get("ADMIN_PASSWORD_HASH") or "").strip()
    bootstrap_version = (os.environ.get("ADMIN_BOOTSTRAP_VERSION") or "").strip()
    is_production = bool(os.environ.get("RENDER") or os.environ.get("ENVIRONMENT") == "production")

    if not admin_email:
        if is_production:
            raise RuntimeError("ADMIN_EMAIL deve ser configurado em produção")
        admin_email = "admin@example.com"

    if not admin_password_hash:
        if not admin_password:
            if is_production:
                raise RuntimeError("ADMIN_PASSWORD ou ADMIN_PASSWORD_HASH deve ser configurado em produção")
            admin_password = "dev-only-admin-password"
        if len(admin_password) < MIN_PASSWORD_LENGTH:
            raise RuntimeError(f"ADMIN_PASSWORD deve ter ao menos {MIN_PASSWORD_LENGTH} caracteres")
        admin_password_hash = hash_password(admin_password)
    elif not admin_password_hash.startswith(("$2a$", "$2b$", "$2y$")):
        raise RuntimeError("ADMIN_PASSWORD_HASH inválido")

    existing = await db.users.find_one({"email": admin_email})
    if existing is None:
        conflicting = await db.users.find_one({"email": admin_email})
        if conflicting is not None:
            raise RuntimeError("ADMIN_EMAIL já pertence a outra conta")

        any_admin = await db.users.find_one({"role": "admin"})
        if any_admin is not None:
            await db.users.update_one(
                {"_id": any_admin["_id"]},
                {"$set": {
                    "email": admin_email,
                    "password_hash": admin_password_hash,
                    "bootstrap_version": bootstrap_version,
                }, "$inc": {"token_version": 1}},
            )
            await db.login_attempts.delete_many({"email": admin_email})
            logger.info("Admin bootstrap credentials synchronized")
            return

        await db.users.insert_one({
            "email": admin_email,
            "password_hash": admin_password_hash,
            "name": "Administrador",
            "role": "admin",
            "token_version": 0,
            "bootstrap_version": bootstrap_version,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        logger.info("Admin seeded")
        return

    if existing.get("role") != "admin":
        raise RuntimeError("ADMIN_EMAIL pertence a uma conta sem privilégio de administrador")

    if bootstrap_version and existing.get("bootstrap_version") != bootstrap_version:
        await db.users.update_one(
            {"_id": existing["_id"]},
            {"$set": {
                "password_hash": admin_password_hash,
                "bootstrap_version": bootstrap_version,
            }, "$inc": {"token_version": 1}},
        )
        await db.login_attempts.delete_many({"email": admin_email})
        logger.info("Admin bootstrap password synchronized")
    else:
        logger.info("Admin account already exists; bootstrap already current")


async def seed_categories():
    if await db.categories.count_documents({}) == 0:
        now = datetime.now(timezone.utc).isoformat()
        for c in DEFAULT_CATEGORIES:
            c["created_at"] = now
            await db.categories.insert_one(dict(c))
        logger.info("Categories seeded")


@app.on_event("startup")
async def startup():
    try:
        database_info = await client.admin.command("ping")
        if DB_ENGINE == "sqlserver":
            logger.info(
                "SQL Server conectado: %s (versão %s)",
                db_name,
                database_info.get("version", "desconhecida"),
            )
        else:
            logger.info("MongoDB conectado: %s", db_name)
    except Exception as e:
        logger.error("Falha ao conectar no banco %s: %s", DB_ENGINE, e)
        raise
    await db.users.create_index("email", unique=True)
    await db.login_attempts.create_index("identifier")
    await db.login_attempts.create_index("expires_at", expireAfterSeconds=0)
    await db.rate_limits.create_index("expires_at", expireAfterSeconds=0)
    await db.tickets.create_index("ticket_number", unique=True)
    await db.tickets.create_index([("category_id", 1), ("created_at", -1)])
    await db.tickets.create_index([("category_ids", 1), ("created_at", -1)])
    await db.tickets.create_index([("status", 1), ("created_at", -1)])
    await db.tickets.create_index("requester.email")
    await db.tickets.create_index("requester.matricula")
    await db.password_reset_tokens.create_index("token_hash", unique=True)
    await db.password_reset_tokens.create_index("email")
    await db.password_reset_tokens.create_index("purge_at", expireAfterSeconds=0)
    await db.password_reset_requests.create_index("email")
    await db.password_reset_requests.create_index("expires_at", expireAfterSeconds=0)
    await db.audit_logs.create_index([("created_at", -1)])
    await db.audit_logs.create_index([("actor_id", 1), ("created_at", -1)])
    await db.email_events.create_index([("email", 1), ("received_at", -1)])
    await db.email_events.create_index([("message_id", 1), ("received_at", -1)])
    await db.users.create_index([("last_login_at", -1)])
    await seed_admin()
    await seed_categories()
    logger.info("Email delivery configured: %s", "yes" if email_delivery_configured() else "no")
    try:
        await init_storage()
        logger.info("Storage inicializado (%s)", "Emergent" if EMERGENT_KEY else f"local: {LOCAL_STORAGE_DIR}")
    except Exception as e:
        logger.warning("Storage remoto indisponível, usando disco local: %s", e)


@api_router.get("/")
async def root():
    return {"message": "Gestão de Materiais API", "status": "ok"}


@api_router.get("/health")
async def health():
    try:
        database_info = await client.admin.command("ping")
        response = {
            "status": "ok",
            "database": "connected",
            "database_engine": DB_ENGINE,
        }
        if DB_ENGINE == "sqlserver" and database_info.get("version"):
            response["database_version"] = database_info["version"]
        return response
    except Exception as e:
        logger.error("Health check falhou (%s): %s", DB_ENGINE, e)
        raise HTTPException(status_code=503, detail="Serviço temporariamente indisponível")


app.include_router(api_router)

configured_origins = [
    origin.strip().rstrip("/")
    for origin in os.environ.get("CORS_ORIGINS", "").split(",")
    if origin.strip()
]
allowed_origins = list(dict.fromkeys([
    FRONTEND_URL,
    "http://localhost:3000",
    "http://127.0.0.1:3000",
    *configured_origins,
]))


@app.middleware("http")
async def security_headers_and_origin_guard(request: Request, call_next):
    unsafe_method = request.method.upper() in {"POST", "PUT", "PATCH", "DELETE"}
    has_session_cookie = bool(request.cookies.get("access_token") or request.cookies.get("refresh_token"))
    if request.url.path.startswith("/api/") and unsafe_method and has_session_cookie:
        origin = (request.headers.get("origin") or "").rstrip("/")
        if origin and origin not in allowed_origins:
            return JSONResponse(status_code=403, content={"detail": "Origem da requisição não permitida"})

    response = await call_next(request)
    if request.url.path.startswith("/api/auth/") or request.url.path.startswith("/admin"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
    return response


app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=allowed_origins,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.add_middleware(GZipMiddleware, minimum_size=1000, compresslevel=5)

FRONTEND_BUILD_DIR = (ROOT_DIR.parent / "frontend" / "build").resolve()

if FRONTEND_BUILD_DIR.exists():
    @app.get("/{full_path:path}", include_in_schema=False)
    async def serve_spa(full_path: str):
        requested = (FRONTEND_BUILD_DIR / full_path).resolve()
        if FRONTEND_BUILD_DIR in requested.parents and requested.is_file():
            response = FileResponse(requested)
            if "/static/" in requested.as_posix():
                response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
            elif requested.name == "index.html":
                response.headers["Cache-Control"] = "no-cache"
            else:
                response.headers["Cache-Control"] = "public, max-age=3600"
            return response
        response = FileResponse(FRONTEND_BUILD_DIR / "index.html")
        response.headers["Cache-Control"] = "no-cache"
        return response


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
